#!/usr/bin/env python3
"""Build a gated Chinese feasibility-report DOCX from assembled Markdown."""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from collections import Counter
from copy import deepcopy
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor

from build_policy_section_material import build_material as build_policy_material
from knowledge_db import apply_migrations, connect, sha256_file, sha256_text


USABLE_WIDTH_DXA = 8240
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
INLINE_PATTERN = re.compile(
    r"(\*\*.+?\*\*|(?<!\*)\*[^*\n]+?\*(?!\*)|`[^`\n]+`|\[[^\]\n]+\]\([^)\n]+\))"
)


def resolve_format_authority(
    template: Path | None,
    format_config: Path | None,
) -> tuple[Path | None, dict[str, Any]]:
    if format_config is None:
        return (template.resolve() if template else None), {}
    config_path = format_config.resolve()
    if not config_path.is_file():
        raise FileNotFoundError(config_path)
    payload = json.loads(config_path.read_text(encoding="utf-8-sig"))
    if payload.get("schema_version") != "1.0":
        raise ValueError("unsupported Word format authority schema_version")
    if payload.get("status") != "confirmed":
        raise ValueError("Word format authority must be confirmed before use")
    if not str(payload.get("confirmed_by") or "").strip() or not str(
        payload.get("confirmed_at") or ""
    ).strip():
        raise ValueError("confirmed Word format authority requires confirmed_by and confirmed_at")
    authority = payload.get("authority")
    if not isinstance(authority, dict) or not authority.get("template_path"):
        raise ValueError("Word format authority requires authority.template_path")
    configured_template = Path(str(authority["template_path"]))
    if not configured_template.is_absolute():
        configured_template = config_path.parent / configured_template
    configured_template = configured_template.resolve()
    if not configured_template.is_file():
        raise FileNotFoundError(configured_template)
    if template is not None and template.resolve() != configured_template:
        raise ValueError("explicit Word template conflicts with the confirmed format authority")
    expected_hash = str(authority.get("template_sha256") or "").lower()
    actual_hash = sha256_file(configured_template)
    if not expected_hash or actual_hash.lower() != expected_hash:
        raise ValueError("confirmed Word template hash does not match the format authority")
    semantic_styles = payload.get("semantic_styles") or {}
    if not isinstance(semantic_styles, dict):
        raise ValueError("semantic_styles must be an object")
    evidence = payload.get("evidence") or {}
    if not isinstance(evidence, dict):
        raise ValueError("evidence must be an object")
    resolved_evidence: dict[str, str] = {}
    for label, path_key, hash_key in (
        ("format profile", "format_profile_path", "format_profile_sha256"),
        ("style contract", "style_contract_path", "style_contract_sha256"),
    ):
        configured_path = str(evidence.get(path_key) or "").strip()
        if not configured_path:
            continue
        evidence_path = Path(configured_path)
        if not evidence_path.is_absolute():
            evidence_path = config_path.parent / evidence_path
        evidence_path = evidence_path.resolve()
        if not evidence_path.is_file():
            raise FileNotFoundError(evidence_path)
        expected_evidence_hash = str(evidence.get(hash_key) or "").lower()
        if not expected_evidence_hash or sha256_file(evidence_path).lower() != expected_evidence_hash:
            raise ValueError(f"confirmed {label} hash does not match the format authority")
        resolved_evidence[path_key] = str(evidence_path)
    payload["_config_path"] = str(config_path)
    payload["_config_sha256"] = sha256_file(config_path)
    payload["_resolved_template"] = str(configured_template)
    payload["_resolved_evidence"] = resolved_evidence
    return configured_template, payload


def _style_name(document: Document, mapping: dict[str, str], key: str, default: str) -> str:
    name = str(mapping.get(key) or default)
    if name not in {style.name for style in document.styles}:
        raise ValueError(f"Word format authority references a missing style: {key}={name}")
    return name


def set_run_font(run, east_asia: str, size: float, *, bold: bool = False) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")


def set_style_font(style, east_asia: str, size: float, *, bold: bool = False) -> None:
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().rFonts
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")


def _append_val(parent, tag: str, value: str) -> OxmlElement:
    element = OxmlElement(tag)
    element.set(qn("w:val"), value)
    parent.append(element)
    return element


def _insert_num_pr(paragraph_properties, num_pr: OxmlElement) -> None:
    existing = paragraph_properties.find(qn("w:numPr"))
    if existing is not None:
        paragraph_properties.remove(existing)
    paragraph_properties.insert_element_before(
        num_pr,
        "w:suppressLineNumbers",
        "w:pBdr",
        "w:shd",
        "w:tabs",
        "w:suppressAutoHyphens",
        "w:kinsoku",
        "w:wordWrap",
        "w:overflowPunct",
        "w:topLinePunct",
        "w:autoSpaceDE",
        "w:autoSpaceDN",
        "w:bidi",
        "w:adjustRightInd",
        "w:snapToGrid",
        "w:spacing",
        "w:ind",
        "w:contextualSpacing",
        "w:mirrorIndents",
        "w:suppressOverlap",
        "w:jc",
        "w:textDirection",
        "w:textAlignment",
        "w:textboxTightWrap",
        "w:outlineLvl",
        "w:divId",
        "w:cnfStyle",
        "w:rPr",
        "w:sectPr",
        "w:pPrChange",
    )


def _next_numbering_id(numbering, element_name: str, attribute_name: str) -> int:
    values = []
    for element in numbering.findall(qn(element_name)):
        raw = element.get(qn(attribute_name))
        if raw is not None and raw.isdigit():
            values.append(int(raw))
    return max(values, default=0) + 1


def configure_default_heading_numbering(document: Document) -> int:
    """Create deterministic H1-H7 numbering for the no-template working preset."""

    numbering = document.part.numbering_part.element
    abstract_num_id = _next_numbering_id(
        numbering, "w:abstractNum", "w:abstractNumId"
    )
    num_id = _next_numbering_id(numbering, "w:num", "w:numId")

    abstract_num = OxmlElement("w:abstractNum")
    abstract_num.set(qn("w:abstractNumId"), str(abstract_num_id))
    _append_val(abstract_num, "w:nsid", "4D454449")
    _append_val(abstract_num, "w:multiLevelType", "multilevel")
    _append_val(abstract_num, "w:tmpl", "46454153")

    for level in range(7):
        heading_style = document.styles[f"Heading {level + 1}"]
        item = OxmlElement("w:lvl")
        item.set(qn("w:ilvl"), str(level))
        _append_val(item, "w:start", "1")
        _append_val(
            item,
            "w:numFmt",
            "chineseCountingThousand" if level == 0 else "decimal",
        )
        _append_val(item, "w:pStyle", heading_style.style_id)
        if level > 0:
            item.append(OxmlElement("w:isLgl"))
        _append_val(item, "w:suff", "space")
        level_text = "第%1章" if level == 0 else ".".join(
            f"%{index}" for index in range(1, level + 2)
        ) + "."
        _append_val(item, "w:lvlText", level_text)
        _append_val(item, "w:lvlJc", "left")
        item_ppr = OxmlElement("w:pPr")
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), "0")
        indent.set(qn("w:firstLine"), "0")
        item_ppr.append(indent)
        item.append(item_ppr)
        abstract_num.append(item)

        style_num_pr = OxmlElement("w:numPr")
        _append_val(style_num_pr, "w:ilvl", str(level))
        _append_val(style_num_pr, "w:numId", str(num_id))
        _insert_num_pr(heading_style.element.get_or_add_pPr(), style_num_pr)

    numbering.append(abstract_num)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    _append_val(num, "w:abstractNumId", str(abstract_num_id))
    numbering.append(num)
    return num_id


def bind_heading_numbering(paragraph, style) -> bool:
    """Copy effective style numbering to the paragraph for Word/WPS parity."""

    style_ppr = style.element.find(qn("w:pPr"))
    style_num_pr = style_ppr.find(qn("w:numPr")) if style_ppr is not None else None
    if style_num_pr is None:
        return False
    num_id = style_num_pr.find(qn("w:numId"))
    ilvl = style_num_pr.find(qn("w:ilvl"))
    if num_id is None or not str(num_id.get(qn("w:val")) or "").isdigit():
        return False
    if int(num_id.get(qn("w:val"))) <= 0 or ilvl is None:
        return False
    _insert_num_pr(paragraph._p.get_or_add_pPr(), deepcopy(style_num_pr))
    return True


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.175)
    section.right_margin = Cm(3.175)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)

    normal = document.styles["Normal"]
    set_style_font(normal, "仿宋_GB2312", 14)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Pt(28)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    heading_tokens = {
        "Heading 1": ("华文中宋", 22, 18, 12, WD_ALIGN_PARAGRAPH.CENTER),
        "Heading 2": ("黑体", 18, 14, 8, WD_ALIGN_PARAGRAPH.LEFT),
        "Heading 3": ("黑体", 16, 12, 6, WD_ALIGN_PARAGRAPH.LEFT),
        "Heading 4": ("楷体_GB2312", 16, 10, 4, WD_ALIGN_PARAGRAPH.LEFT),
        "Heading 5": ("仿宋_GB2312", 14, 8, 4, WD_ALIGN_PARAGRAPH.LEFT),
        "Heading 6": ("仿宋_GB2312", 14, 6, 3, WD_ALIGN_PARAGRAPH.LEFT),
        "Heading 7": ("仿宋_GB2312", 14, 4, 2, WD_ALIGN_PARAGRAPH.LEFT),
    }
    for style_name, (font, size, before, after, alignment) in heading_tokens.items():
        style = document.styles[style_name]
        set_style_font(style, font, size, bold=style_name != "Heading 1")
        style.paragraph_format.alignment = alignment
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.0
    document.styles["Heading 1"].paragraph_format.page_break_before = True

    for style_name in ("List Bullet", "List Number"):
        style = document.styles[style_name]
        set_style_font(style, "仿宋_GB2312", 14)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    configure_default_heading_numbering(document)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)


def configure_header_footer(document: Document, project_name: str) -> None:
    for section in document.sections:
        paragraph = section.header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(project_name)
        set_run_font(run, "宋体", 9)
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(paragraph, "PAGE", "1")


def clear_container(element) -> None:
    for child in list(element):
        element.remove(child)
    element.append(OxmlElement("w:p"))


def dominant_template_geometry(document: Document) -> dict[str, Any]:
    profiles = []
    for index, section in enumerate(document.sections):
        profile = {
            "page_width": section.page_width,
            "page_height": section.page_height,
            "top_margin": section.top_margin,
            "bottom_margin": section.bottom_margin,
            "left_margin": section.left_margin,
            "right_margin": section.right_margin,
            "header_distance": section.header_distance,
            "footer_distance": section.footer_distance,
            "gutter": section.gutter,
        }
        key = tuple(int(profile[name] or 0) for name in profile)
        profiles.append((index, profile, key, int(section.page_width) <= int(section.page_height)))
    portrait_profiles = [item for item in profiles if item[3]]
    candidates = portrait_profiles or profiles
    counts = Counter(item[2] for item in candidates)
    best_key = max(counts, key=lambda key: (counts[key], -next(item[0] for item in candidates if item[2] == key)))
    return next(item[1] for item in candidates if item[2] == best_key)


def scrub_template(document: Document) -> None:
    geometry = dominant_template_geometry(document)
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    removable_relationships = {
        RT.IMAGE, RT.HYPERLINK, RT.OLE_OBJECT, RT.PACKAGE,
        RT.COMMENTS, RT.CUSTOM_XML,
    }
    parts = [document.part]
    parts.extend(section.header.part for section in document.sections)
    parts.extend(section.footer.part for section in document.sections)
    for part in parts:
        for relationship_id, relationship in list(part.rels.items()):
            if relationship.reltype in removable_relationships:
                part.drop_rel(relationship_id)
    for section in document.sections:
        clear_container(section.header._element)
        clear_container(section.footer._element)
    target = document.sections[0]
    for name, value in geometry.items():
        setattr(target, name, value)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, column_widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_width = table._tbl.tblPr.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), str(sum(column_widths)))
    table_width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in column_widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Cm(column_widths[index] / 567)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_width.set(qn("w:w"), str(column_widths[index]))
            tc_width.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    header = row_properties.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        row_properties.append(header)
    header.set(qn("w:val"), "true")


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline_markdown(paragraph, value: str) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(value):
        if match.start() > position:
            paragraph.add_run(value[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        elif token.startswith("["):
            link = re.fullmatch(r"\[([^]]+)]\(([^)]+)\)", token)
            if link:
                add_hyperlink(paragraph, link.group(1), link.group(2))
        position = match.end()
    if position < len(value):
        paragraph.add_run(value[position:])


def plain_inline_text(value: str) -> str:
    value = re.sub(r"\[([^]]+)]\([^)]+\)", r"\1", value)
    return re.sub(r"\*\*|`|(?<!\*)\*(?!\*)", "", value).strip()


def add_markdown_table(
    document: Document,
    rows: list[list[str]],
    semantic_styles: dict[str, str] | None = None,
) -> None:
    semantic_styles = semantic_styles or {}
    column_count = max(len(row) for row in rows)
    normalized = [row + [""] * (column_count - len(row)) for row in rows]
    if len(normalized) > 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in normalized[1]):
        normalized.pop(1)
    table = document.add_table(rows=len(normalized), cols=column_count)
    table.style = _style_name(document, semantic_styles, "table", "Table Grid")
    base = USABLE_WIDTH_DXA // column_count
    widths = [base] * column_count
    widths[-1] += USABLE_WIDTH_DXA - sum(widths)
    for row_index, row in enumerate(normalized):
        for column_index, value in enumerate(row):
            paragraph = table.cell(row_index, column_index).paragraphs[0]
            paragraph.style = _style_name(
                document,
                semantic_styles,
                "table_header" if row_index == 0 else "table_body",
                "Normal",
            )
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline_markdown(paragraph, value)
            if not semantic_styles:
                for run in paragraph.runs:
                    set_run_font(run, "黑体" if row_index == 0 else "宋体", 10.5, bold=row_index == 0 or bool(run.bold))
    set_table_geometry(table, widths)
    if table.rows:
        set_repeat_table_header(table.rows[0])


def parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def current_adopted_hash(conn, project_id: str) -> str:
    rows = conn.execute(
        """
        SELECT d.plan_id,d.draft_version_id,d.version_no,d.content
        FROM draft_section_version d
        JOIN section_composition_plan p ON p.plan_id=d.plan_id
        WHERE p.project_id=? AND d.status='adopted'
          AND p.applicability_status<>'not_applicable'
          AND p.version_no=(
            SELECT MAX(p2.version_no) FROM section_composition_plan p2
            WHERE p2.project_id=p.project_id AND p2.chapter_code=p.chapter_code
              AND p2.applicability_status<>'not_applicable'
          )
        ORDER BY d.plan_id
        """,
        (project_id,),
    ).fetchall()
    manifest = [
        {
            "plan_id": row["plan_id"],
            "draft_version_id": row["draft_version_id"],
            "version_no": row["version_no"],
            "content_sha256": sha256_text(row["content"]),
        }
        for row in rows
    ]
    return sha256_text(json.dumps(manifest, ensure_ascii=False, sort_keys=True, separators=(",", ":")))


def verify_policy_validation_freshness(
    summary: dict[str, Any],
    current_material: dict[str, Any] | None,
    *,
    has_policy_chapters: bool,
) -> None:
    if not has_policy_chapters:
        return
    if current_material is None:
        raise RuntimeError("formal delivery blocked: current policy material is unavailable")
    delivery_blockers = list(
        current_material.get("quality", {}).get("delivery_blockers", [])
    )
    if current_material.get("delivery_eligible") is not True or delivery_blockers:
        detail = "、".join(delivery_blockers[:10]) or "delivery_eligible=false"
        raise RuntimeError(
            "formal delivery blocked: current policy material is not delivery eligible: "
            + detail
        )
    if (
        summary.get("policy_match_run_id") != current_material.get("match_run_id")
        or summary.get("policy_material_signature")
        != current_material.get("material_signature")
    ):
        raise RuntimeError("formal delivery blocked: policy validation is stale")


def authorize_delivery(
    markdown: str,
    template: Path | None,
    database: Path | None,
    project_code: str,
) -> dict[str, str]:
    if template is None:
        raise RuntimeError("formal delivery blocked: a confirmed Word template is required")
    if database is None or not project_code.strip():
        raise RuntimeError("formal delivery blocked: --database and --project-code are required")
    from assemble_report_markdown import assemble

    assembled = assemble(database, project_code, mode="delivery")
    if sha256_text(assembled["content"]) != sha256_text(markdown):
        raise RuntimeError("formal delivery blocked: Markdown is not the current database assembly")
    with connect(database.resolve()) as conn:
        apply_migrations(conn)
        project = conn.execute(
            "SELECT * FROM project WHERE project_code=?", (project_code,)
        ).fetchone()
        if project is None:
            raise RuntimeError(f"project_code {project_code} is not initialized")
        if not project["owner_name"] or "【待" in project["official_name"]:
            raise RuntimeError("formal delivery blocked: official project name and owner must be confirmed")
        scope_count = conn.execute(
            "SELECT COUNT(*) FROM project_scope_item WHERE project_id=? AND customer_scope=1",
            (project["project_id"],),
        ).fetchone()[0]
        if scope_count:
            confirmed_baseline = conn.execute(
                "SELECT 1 FROM scope_baseline WHERE project_id=? AND status='confirmed' LIMIT 1",
                (project["project_id"],),
            ).fetchone()
            if not confirmed_baseline:
                raise RuntimeError("formal delivery blocked: scope baseline is not confirmed")
        validation = conn.execute(
            """
            SELECT * FROM validation_run WHERE project_id=?
              AND validation_type='full_report_delivery'
            ORDER BY completed_at DESC,validation_run_id DESC LIMIT 1
            """,
            (project["project_id"],),
        ).fetchone()
        if validation is None or validation["status"] != "passed" or validation["blocking_count"]:
            raise RuntimeError("formal delivery blocked: latest delivery validation did not pass")
        summary = json.loads(validation["summary_json"] or "{}")
        content_sha256 = current_adopted_hash(conn, project["project_id"])
        if summary.get("content_sha256") != content_sha256:
            raise RuntimeError("formal delivery blocked: delivery validation is stale")
        if summary.get("outline_hash") != assembled["outline_hash"]:
            raise RuntimeError("formal delivery blocked: confirmed outline validation is stale")
        active_policy_chapter_count = int(
            conn.execute(
                """
                SELECT COUNT(DISTINCT chapter_code) FROM section_composition_plan
                WHERE project_id=? AND chapter_code IN ('1.2.1','2.1.1')
                  AND applicability_status<>'not_applicable'
                """,
                (project["project_id"],),
            ).fetchone()[0]
        )
        if project["document_type"] == "feasibility_study" and active_policy_chapter_count != 2:
            raise RuntimeError(
                "formal delivery blocked: feasibility study requires both 1.2.1 and 2.1.1 policy chapters"
            )
        has_policy_chapters = active_policy_chapter_count > 0
        current_policy_material = None
        if has_policy_chapters:
            try:
                current_policy_material = build_policy_material(
                    database, project_code, mode="delivery"
                )
            except ValueError as exc:
                raise RuntimeError(
                    f"formal delivery blocked: current policy material is unavailable: {exc}"
                ) from exc
        verify_policy_validation_freshness(
            summary,
            current_policy_material,
            has_policy_chapters=has_policy_chapters,
        )
    return {
        "validation_run_id": validation["validation_run_id"],
        "validated_content_sha256": content_sha256,
        "outline_version_id": assembled["outline_version_id"],
        "outline_hash": assembled["outline_hash"],
        "policy_match_run_id": summary.get("policy_match_run_id", ""),
        "policy_material_signature": summary.get("policy_material_signature", ""),
    }


def audit_docx_markdown_residue(path: Path) -> dict[str, int]:
    with zipfile.ZipFile(path) as package:
        root = ET.fromstring(package.read("word/document.xml"))
    paragraphs = []
    for paragraph in root.findall(f".//{{{W_NS}}}p"):
        text = "".join(node.text or "" for node in paragraph.findall(f".//{{{W_NS}}}t"))
        paragraphs.append(text)
    return {
        "heading_markers": sum(bool(re.match(r"^#{1,7}\s+", text)) for text in paragraphs),
        "bold_markers": sum(text.count("**") for text in paragraphs),
        "code_markers": sum(text.count("`") for text in paragraphs),
    }


def build_docx(
    markdown: str,
    output: Path,
    project_name: str,
    owner_name: str = "",
    template: Path | None = None,
    *,
    mode: str = "working",
    database: Path | None = None,
    project_code: str = "",
    format_config: Path | None = None,
) -> dict[str, Any]:
    if mode not in {"working", "delivery"}:
        raise ValueError("mode must be working or delivery")
    template, format_authority = resolve_format_authority(template, format_config)
    if template and not template.is_file():
        raise FileNotFoundError(template)
    authorization = authorize_delivery(markdown, template, database, project_code) if mode == "delivery" else {}
    markdown = re.sub(r"<!--[\s\S]*?-->", "", markdown)

    if template:
        document = Document(template)
        scrub_template(document)
        preset = (
            "confirmed format authority, template styles and page geometry"
            if format_authority
            else "confirmed-template styles and page geometry"
        )
    else:
        document = Document()
        configure_document(document)
        preset = "A4 Chinese government feasibility working preset"
    configure_header_footer(document, project_name)
    semantic_styles = {
        str(key): str(value)
        for key, value in (format_authority.get("semantic_styles") or {}).items()
    }
    heading_style_names = {
        level: _style_name(
            document,
            semantic_styles,
            f"heading_{level}",
            f"Heading {level}",
        )
        for level in range(1, 8)
    }
    body_style_name = _style_name(document, semantic_styles, "body", "Normal")
    numbered_heading_levels = set()
    for level in range(1, 8):
        style = document.styles[heading_style_names[level]]
        paragraph_properties = style.element.find(qn("w:pPr"))
        if paragraph_properties is not None and paragraph_properties.find(qn("w:numPr")) is not None:
            numbered_heading_levels.add(level)
    document.core_properties.title = f"{project_name}可行性研究报告"
    document.core_properties.subject = "医疗信息化政府投资项目可行性研究报告"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = f"Layout source: {preset}. Template body and personal metadata were scrubbed."

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Cm(5.5)
    title.paragraph_format.space_after = Pt(24)
    set_run_font(title.add_run(project_name), "华文中宋", 22)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(subtitle.add_run("可行性研究报告"), "华文中宋", 22)
    if mode == "working":
        status = document.add_paragraph()
        status.alignment = WD_ALIGN_PARAGRAPH.CENTER
        status.paragraph_format.space_before = Pt(18)
        run = status.add_run("工作稿（未通过正式交付门禁）")
        set_run_font(run, "黑体", 14, bold=True)
        run.font.color.rgb = RGBColor(192, 0, 0)
    if owner_name:
        owner = document.add_paragraph()
        owner.alignment = WD_ALIGN_PARAGRAPH.CENTER
        owner.paragraph_format.space_before = Cm(7 if mode == "working" else 8)
        set_run_font(owner.add_run(owner_name), "宋体", 14)
    document.add_page_break()
    toc_title = document.add_paragraph("目录")
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(toc_title.runs[0], "黑体", 16, bold=True)
    toc = document.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "请在 Word 中更新目录")
    document.add_page_break()

    lines = markdown.splitlines()
    index = 0
    body_started = False
    paragraph_count = 0
    table_count = 0
    heading_levels = {str(level): 0 for level in range(1, 8)}
    explicitly_numbered_headings = 0
    code_fence = False
    while index < len(lines):
        line = lines[index].rstrip()
        if line.strip().startswith("```"):
            code_fence = not code_fence
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        heading = re.match(r"^(#{1,7})\s+(.+)$", line)
        if heading and not code_fence:
            text = plain_inline_text(heading.group(2))
            if not body_started and not re.match(r"^第\d+章", text):
                index += 1
                continue
            body_started = True
            level = len(heading.group(1))
            if level in numbered_heading_levels:
                if level == 1:
                    text = re.sub(r"^第(?:\d+|[一二三四五六七八九十百]+)章\s*", "", text)
                else:
                    text = re.sub(r"^\d+(?:\.\d+){0,6}[.、．]?\s*", "", text)
            paragraph = document.add_paragraph(text, style=heading_style_names[level])
            if level in numbered_heading_levels:
                if not bind_heading_numbering(
                    paragraph, document.styles[heading_style_names[level]]
                ):
                    raise RuntimeError(
                        f"failed to bind multilevel numbering to Heading {level}"
                    )
                explicitly_numbered_headings += 1
            heading_levels[str(level)] += 1
            index += 1
            continue
        if not body_started:
            index += 1
            continue
        if re.fullmatch(r"\s*(?:---+|___+|\*\*\*+)\s*", line):
            index += 1
            continue
        if not code_fence and line.strip().startswith("|") and "|" in line.strip()[1:]:
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(parse_table_row(lines[index]))
                index += 1
            add_markdown_table(document, table_lines, semantic_styles)
            table_count += 1
            continue
        list_item = re.match(r"^(\s*)([-+*]|\d+[.)])\s+(.+)$", line)
        if list_item and not code_fence:
            marker = list_item.group(2)
            list_style = (
                _style_name(document, semantic_styles, "list_bullet", "List Bullet")
                if not marker[0].isdigit()
                else _style_name(document, semantic_styles, "list_number", "List Number")
            )
            paragraph = document.add_paragraph(style=list_style)
            depth = min(4, len(list_item.group(1).replace("\t", "    ")) // 2)
            paragraph.paragraph_format.left_indent = Cm(0.74 + depth * 0.74)
            paragraph.paragraph_format.first_line_indent = Cm(-0.37)
            add_inline_markdown(paragraph, list_item.group(3))
            paragraph_count += 1
            index += 1
            continue
        paragraph = document.add_paragraph(style=body_style_name)
        if code_fence:
            paragraph.paragraph_format.first_line_indent = Pt(0)
            run = paragraph.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(10.5)
        else:
            if line.lstrip().startswith(">"):
                line = line.lstrip()[1:].lstrip()
                paragraph.paragraph_format.left_indent = Cm(0.74)
                paragraph.paragraph_format.first_line_indent = Pt(0)
            add_inline_markdown(paragraph, line.strip())
        paragraph_count += 1
        index += 1

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    output = output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    expected_numbered_headings = sum(
        heading_levels[str(level)] for level in numbered_heading_levels
    )
    if explicitly_numbered_headings != expected_numbered_headings:
        raise RuntimeError(
            "generated heading numbering audit failed: "
            f"expected {expected_numbered_headings}, "
            f"bound {explicitly_numbered_headings}"
        )
    document.save(output)
    residue = audit_docx_markdown_residue(output)
    if mode == "delivery" and any(residue.values()):
        output.unlink(missing_ok=True)
        raise RuntimeError(f"formal delivery blocked: Markdown residue remains in DOCX: {residue}")
    output_sha256 = sha256_file(output)
    return {
        "output": str(output),
        "output_sha256": output_sha256,
        "mode": mode,
        "preset": preset,
        "template": str(template) if template else "",
        "template_sha256": sha256_file(template) if template else "",
        "format_config": format_authority.get("_config_path", ""),
        "format_config_sha256": format_authority.get("_config_sha256", ""),
        "format_profile_id": format_authority.get("profile_id", ""),
        "semantic_styles": semantic_styles,
        "heading_levels": heading_levels,
        "numbered_heading_levels": sorted(numbered_heading_levels),
        "numbered_heading_paragraphs": explicitly_numbered_headings,
        "heading_numbering_audit": {
            "status": "pass",
            "expected": expected_numbered_headings,
            "explicitly_bound": explicitly_numbered_headings,
        },
        "headings": sum(heading_levels.values()),
        "paragraphs": paragraph_count,
        "tables": table_count,
        "markdown_residue": residue,
        "authorization": authorization,
        "render_required": True,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--project-name", required=True)
    parser.add_argument("--owner-name", default="")
    parser.add_argument("--template", type=Path)
    parser.add_argument("--format-config", type=Path)
    parser.add_argument("--mode", choices=("working", "delivery"), default="working")
    parser.add_argument("--database", type=Path)
    parser.add_argument("--project-code", default="")
    parser.add_argument("--summary", type=Path)
    args = parser.parse_args()
    result = build_docx(
        args.input.read_text(encoding="utf-8-sig"), args.output,
        args.project_name, args.owner_name, args.template,
        mode=args.mode, database=args.database, project_code=args.project_code,
        format_config=args.format_config,
    )
    text = json.dumps(result, ensure_ascii=False, indent=2)
    if args.summary:
        args.summary.parent.mkdir(parents=True, exist_ok=True)
        args.summary.write_text(text + "\n", encoding="utf-8")
    print(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
