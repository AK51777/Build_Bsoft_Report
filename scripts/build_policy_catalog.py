#!/usr/bin/env python3
"""Build a reviewable policy-catalog package from the department XLSX index."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from knowledge_db import now_iso, stable_id


REQUIRED_HEADERS = {"级别", "类别", "字号", "标题", "发文日期", "发布部门与机构", "索引号", "外部链接"}
GROUP_PATTERN = re.compile(r"\((CN(?:\d+(?:-\d+)?)|CNX)\)\s*(.*)", re.IGNORECASE)
KEYWORD_SPLIT_PATTERN = re.compile(r"[/、,，;；\n]+")


def clean_text(value: Any) -> str:
    return re.sub(r"[ \t]+", " ", str(value or "").replace("\u00a0", " ")).strip()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_date(value: Any) -> tuple[str, str]:
    if isinstance(value, datetime):
        return value.date().isoformat(), value.isoformat()
    if isinstance(value, date):
        return value.isoformat(), value.isoformat()
    raw = clean_text(value)
    for pattern in (r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})$", r"^(\d{4})(\d{2})(\d{2})$"):
        match = re.match(pattern, raw)
        if match:
            try:
                parsed = date(*(int(part) for part in match.groups()))
            except ValueError:
                break
            return parsed.isoformat(), raw
    return "", raw


def keyword_tags(value: Any) -> list[str]:
    tags = [clean_text(item) for item in KEYWORD_SPLIT_PATTERN.split(clean_text(value))]
    return list(dict.fromkeys(tag for tag in tags if tag and tag != "-"))


def workbook_sheet(workbook, requested_sheet: str = ""):
    candidates = []
    for sheet in workbook.worksheets:
        if requested_sheet and sheet.title != requested_sheet:
            continue
        for row_number in range(1, min(sheet.max_row, 15) + 1):
            headers = [clean_text(sheet.cell(row_number, column).value) for column in range(1, sheet.max_column + 1)]
            if REQUIRED_HEADERS.issubset(set(headers)):
                candidates.append((sheet.max_row, sheet, row_number, headers))
                break
    if not candidates:
        label = requested_sheet or "a visible sheet with the required policy headers"
        raise ValueError(f"policy worksheet not found: {label}")
    candidates.sort(key=lambda item: (item[1].sheet_state == "visible", item[0]), reverse=True)
    return candidates[0][1:]


def build_catalog(path: Path, *, sheet_name: str = "", title: str = "") -> dict[str, Any]:
    path = path.resolve()
    if not path.is_file() or path.suffix.lower() != ".xlsx":
        raise FileNotFoundError(f"policy catalog XLSX not found: {path}")
    workbook = load_workbook(path, read_only=False, data_only=True)
    sheet, header_row, headers = workbook_sheet(workbook, sheet_name)
    columns = {header: index + 1 for index, header in enumerate(headers) if header}
    file_hash = sha256_file(path)
    catalog_id = stable_id("POLICYCATALOG", file_hash, sheet.title)
    current_group_code = ""
    current_group_name = ""
    records: list[dict[str, Any]] = []
    for row_number in range(header_row + 1, sheet.max_row + 1):
        level = clean_text(sheet.cell(row_number, columns["级别"]).value)
        title_value = clean_text(sheet.cell(row_number, columns["标题"]).value)
        index_no = clean_text(sheet.cell(row_number, columns["索引号"]).value)
        if level and not title_value:
            group_match = GROUP_PATTERN.search(level)
            if group_match:
                current_group_code = group_match.group(1).upper()
                current_group_name = clean_text(group_match.group(2))
            continue
        if not title_value or not index_no:
            continue
        document_no = clean_text(sheet.cell(row_number, columns["字号"]).value)
        issuer = clean_text(sheet.cell(row_number, columns["发布部门与机构"]).value)
        publish_date, publish_date_raw = parse_date(
            sheet.cell(row_number, columns["发文日期"]).value
        )
        external_cell = sheet.cell(row_number, columns["外部链接"])
        external_url = clean_text(external_cell.value)
        if not external_url and external_cell.hyperlink:
            external_url = clean_text(external_cell.hyperlink.target)
        keyword_text = clean_text(sheet.cell(row_number, columns.get("关键词", 0)).value) if columns.get("关键词") else ""
        category_name = clean_text(sheet.cell(row_number, columns["类别"]).value)
        notes = clean_text(sheet.cell(row_number, columns.get("备注", 0)).value) if columns.get("备注") else ""
        file_count_value = sheet.cell(row_number, columns.get("文件数", 0)).value if columns.get("文件数") else None
        file_count = int(file_count_value) if isinstance(file_count_value, (int, float)) else None
        identity_key = stable_id("POLICYIDENTITY", title_value, document_no, issuer)
        entry = {
            "source_row": row_number,
            "source_index_no": index_no,
            "identity_key": identity_key,
            "catalog_group_code": current_group_code,
            "catalog_group_name": current_group_name,
            "authority_level_label": level,
            "category_name": category_name,
            "keyword_text": keyword_text,
            "keyword_tags": keyword_tags(keyword_text),
            "document_no": "" if document_no == "-" else document_no,
            "title": title_value,
            "publish_date": publish_date,
            "publish_date_raw": publish_date_raw,
            "issuer": issuer,
            "file_count": file_count,
            "notes": notes,
            "external_url": external_url,
            "verification_status": "unverified",
            "entry_status": "active",
        }
        records.append(entry)
    if not records:
        raise ValueError("policy catalog does not contain any usable records")
    index_counts = Counter(record["source_index_no"] for record in records)
    index_occurrences: defaultdict[str, int] = defaultdict(int)
    for entry in records:
        index_no = entry["source_index_no"]
        index_occurrences[index_no] += 1
        entry["index_occurrence"] = index_occurrences[index_no]
        entry["index_conflict"] = index_counts[index_no] > 1
        entry["catalog_entry_id"] = stable_id(
            "POLICYCATENTRY", catalog_id, index_no, entry["source_row"]
        )
        entry["row_hash"] = hashlib.sha256(
            json.dumps(entry, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
        ).hexdigest()
    duplicate_groups = [count for count in index_counts.values() if count > 1]
    content = {
        "schema_version": "1.0",
        "catalog_id": catalog_id,
        "catalog_scope": "medical_health_national",
        "title": title or path.stem,
        "permission_scope": "internal_company_reference",
        "source_file": {"file_name": path.name, "sha256": file_hash},
        "worksheet_name": sheet.title,
        "quality_summary": {
            "duplicate_index_groups": len(duplicate_groups),
            "duplicate_index_rows": sum(duplicate_groups),
        },
        "records": records,
    }
    content_hash = hashlib.sha256(
        json.dumps(content, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    payload = {**content, "built_at": now_iso(), "content_hash": content_hash}
    return payload


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--sheet", default="")
    parser.add_argument("--title", default="")
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    payload = build_catalog(args.xlsx, sheet_name=args.sheet, title=args.title)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {
                "catalog_id": payload["catalog_id"],
                "worksheet_name": payload["worksheet_name"],
                "record_count": len(payload["records"]),
                "output": str(args.output.resolve()),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
