from __future__ import annotations

import json
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from apply_scope_capability_decisions import apply_decisions  # noqa: E402
from assemble_report_markdown import assemble  # noqa: E402
from build_evidence_bound_initial_drafts import construction_section  # noqa: E402
from build_section_composition_plan import build_composition_plan  # noqa: E402
from build_standard_knowledge_pack import build_pack  # noqa: E402
from import_standard_knowledge_pack import import_pack  # noqa: E402
from ingest_scope_items_sqlite import ingest_scope_payload  # noqa: E402
from init_project_workbench import initialize_project  # noqa: E402
from knowledge_db import connect  # noqa: E402
from map_scope_capabilities import map_capabilities  # noqa: E402
from report_outline import build_outline_candidate, confirm_outline  # noqa: E402


def blueprint_payload(section_title: str = "项目基本情况") -> dict:
    return {
        "document_type": "feasibility_study",
        "blueprints": [
            {
                "section_role": "project_overview",
                "purpose": "界定项目边界。",
                "required_questions": [],
                "required_fact_categories": [],
                "required_scope_types": [],
                "required_policy_topics": [],
                "required_tables": [],
                "forbidden_content": [],
                "completion_rules": [],
            }
        ],
        "outline": [
            {
                "chapter_code": "1.1.1",
                "section_title": section_title,
                "section_role": "project_overview",
            }
        ],
    }


def make_multi_block_standard(root: Path) -> tuple[Path, Path]:
    solution = root / "multi-block-standard.docx"
    document = Document()
    document.add_heading("建设内容", level=1)
    document.add_heading("电子病历系统", level=2)
    document.add_heading("病历质控", level=3)
    for index in range(1, 6):
        document.add_heading(f"病历质控功能{index}", level=4)
        document.add_paragraph(
            (
                f"病历质控功能{index}围绕规则配置、文书检查、缺陷反馈、整改复核和结果追踪形成完整闭环。"
                "标准方案正文说明使用角色、触发条件、输入输出、异常处理、权限日志和验收证据。"
            )
            * 3
        )
    document.save(solution)

    scope = root / "multi-block-scope.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "主表"
    sheet.append(["大类", "系统名称", "模块名称", "产品功能", "三级医院推荐"])
    sheet.append(["临床业务", "电子病历系统", "病历质控", "病历质控", "推荐"])
    workbook.save(scope)
    return solution, scope


class ReportOutlineAuthorityTests(unittest.TestCase):
    def test_candidate_must_be_confirmed_and_stale_confirmation_cannot_assemble(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            initialized = initialize_project(
                root / "project",
                project_code="OUTLINE-001",
                official_name="目录权威链测试项目",
                owner_name="测试医院",
            )
            database = Path(initialized["database"])
            build_composition_plan(
                database, "OUTLINE-001", blueprint_payload=blueprint_payload()
            )
            candidate = build_outline_candidate(database, "OUTLINE-001")
            self.assertEqual(candidate["status"], "candidate")
            self.assertEqual(
                [node["heading_level"] for node in candidate["nodes"]], [1, 2, 3, 4, 4]
            )
            with self.assertRaisesRegex(RuntimeError, "not confirmed"):
                assemble(database, "OUTLINE-001", mode="working")
            preview = assemble(
                database,
                "OUTLINE-001",
                mode="working",
                allow_candidate_outline=True,
            )
            self.assertEqual(preview["outline_status"], "candidate")
            self.assertIn("目录候选预览", preview["content"])

            edited = json.loads(json.dumps(candidate, ensure_ascii=False))
            section = next(node for node in edited["nodes"] if node["node_kind"] == "section")
            section["title"] = "经确认的项目基本情况"
            fixed_heading = next(
                node for node in edited["nodes"] if node["node_code"] == "1.1.1.1"
            )
            fixed_heading["title"] = "经确认的编制边界"
            confirmed = confirm_outline(
                database,
                "OUTLINE-001",
                outline_version_id=candidate["outline_version_id"],
                confirmed_by="test-reviewer",
                edited_payload=edited,
            )
            self.assertEqual(confirmed["status"], "confirmed")
            with connect(database) as conn:
                conn.execute(
                    """
                    INSERT INTO draft_section_version (
                      draft_version_id,plan_id,version_no,source_type,content,status,
                      check_result_json,created_by,created_at,updated_at
                    ) VALUES (?,?,?,?,?,?,?,?,?,?)
                    """,
                    (
                        "DRAFT-OUTLINE-001",
                        section["source_object_id"],
                        1,
                        "ai",
                        "#### 1.1.1.1 项目定位与编制边界\n\n正文一。\n\n"
                        "#### 1.1.1.2 验收与决策边界\n\n正文二。",
                        "adopted",
                        "{}",
                        "test",
                        "2026-08-17T00:00:00+00:00",
                        "2026-08-17T00:00:00+00:00",
                    ),
                )
                conn.commit()
            assembled = assemble(database, "OUTLINE-001", mode="working")
            self.assertEqual(assembled["outline_status"], "confirmed")
            self.assertIn("### 1.1.1 经确认的项目基本情况", assembled["content"])
            self.assertIn("#### 1.1.1.1 经确认的编制边界", assembled["content"])
            self.assertNotIn("#### 1.1.1.1 项目定位与编制边界", assembled["content"])

            build_composition_plan(
                database,
                "OUTLINE-001",
                blueprint_payload=blueprint_payload("调整后的项目基本情况"),
            )
            replacement = build_outline_candidate(database, "OUTLINE-001")
            self.assertEqual(replacement["status"], "candidate")
            self.assertNotEqual(replacement["source_signature"], confirmed["source_signature"])
            with self.assertRaisesRegex(RuntimeError, "no longer current"):
                confirm_outline(
                    database,
                    "OUTLINE-001",
                    outline_version_id=confirmed["outline_version_id"],
                    confirmed_by="test-reviewer",
                )
            with self.assertRaisesRegex(RuntimeError, "not confirmed"):
                assemble(database, "OUTLINE-001", mode="working")


class FullStandardSolutionAssemblyTests(unittest.TestCase):
    def test_confirmed_capability_carries_every_reviewed_block_in_source_order(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            solution, company_scope = make_multi_block_standard(root)
            pack = build_pack(solution, company_scope, title="全量组装测试知识包")
            capability_payload = pack["capabilities"][0]
            expected_block_ids = capability_payload["standard_block_ids"]
            self.assertGreaterEqual(len(expected_block_ids), 5)
            self.assertEqual(
                capability_payload["block_match_scope"], "capability_or_module_heading"
            )

            initialized = initialize_project(
                root / "project",
                project_code="FULL-BLOCK-001",
                official_name="标准方案全量组装测试项目",
                owner_name="测试医院",
            )
            database = Path(initialized["database"])
            import_pack(database, pack)
            ingest_scope_payload(
                database,
                {
                    "source": {"path": "scope.xlsx", "sha256": "b" * 64},
                    "sheets": [
                        {
                            "name": "建设清单",
                            "detected_columns": {"original_name": ["建设内容"]},
                            "rows": [{"_source_row": 2, "建设内容": "电子病历系统"}],
                        }
                    ],
                },
                project_code="FULL-BLOCK-001",
            )
            conn = sqlite3.connect(database)
            try:
                conn.execute("UPDATE project_scope_item SET status='confirmed'")
                conn.commit()
            finally:
                conn.close()
            mapping = map_capabilities(database, "FULL-BLOCK-001")
            apply_decisions(
                database,
                {
                    "project_code": "FULL-BLOCK-001",
                    "reviewed_by": "test-reviewer",
                    "decisions": [
                        {
                            "map_id": mapping["candidates"][0]["map_id"],
                            "decision": "confirmed",
                            "review_note": "客户范围与能力边界一致",
                        }
                    ],
                },
            )
            result = build_composition_plan(database, "FULL-BLOCK-001")
            construction_result = next(
                item for item in result["plans"] if item["chapter_code"] == "5.1.1"
            )
            self.assertEqual(
                construction_result["full_standard_block_count"], len(expected_block_ids)
            )
            self.assertNotIn(
                "capability_block_scope_review",
                construction_result["missing_source_types"],
            )

            with connect(database) as conn:
                plan = dict(
                    conn.execute(
                        """
                        SELECT p.*,b.section_role FROM section_composition_plan p
                        JOIN section_blueprint b ON b.blueprint_id=p.blueprint_id
                        WHERE p.project_id=(SELECT project_id FROM project WHERE project_code=?)
                          AND p.chapter_code='5.1.1'
                        """,
                        ("FULL-BLOCK-001",),
                    ).fetchone()
                )
                nodes = [
                    dict(row)
                    for row in conn.execute(
                        "SELECT * FROM section_outline_node WHERE plan_id=? ORDER BY ordinal",
                        (plan["plan_id"],),
                    )
                ]
                selected_in_order: list[str] = []
                for node in nodes:
                    metadata = json.loads(node["metadata_json"] or "{}")
                    selected_in_order.extend(metadata.get("block_ids", []))
                self.assertEqual(selected_in_order, expected_block_ids)
                planned = [
                    row[0]
                    for row in conn.execute(
                        """
                        SELECT source_object_id FROM section_plan_source
                        WHERE plan_id=? AND source_type='corpus' AND usage_mode='parameterized'
                        ORDER BY rowid
                        """,
                        (plan["plan_id"],),
                    )
                ]
                self.assertEqual(set(planned), set(expected_block_ids))
                project = dict(
                    conn.execute(
                        "SELECT * FROM project WHERE project_code='FULL-BLOCK-001'"
                    ).fetchone()
                )
                scopes = {
                    row["scope_id"]: dict(row)
                    for row in conn.execute(
                        "SELECT * FROM project_scope_item WHERE project_id=?",
                        (project["project_id"],),
                    )
                }
                capabilities = {
                    row["capability_id"]: dict(row)
                    for row in conn.execute("SELECT * FROM product_capability")
                }
                blocks = {
                    row["block_id"]: dict(row)
                    for row in conn.execute("SELECT * FROM corpus_block")
                }
            for node in nodes:
                node["metadata"] = json.loads(node["metadata_json"] or "{}")
            content = "\n".join(
                construction_section(
                    project,
                    plan,
                    nodes,
                    scopes,
                    capabilities,
                    blocks,
                    ["建设内容总表"],
                )
            )
            for block_id in expected_block_ids:
                self.assertIn(blocks[block_id]["clean_text"], content)


if __name__ == "__main__":
    unittest.main()
