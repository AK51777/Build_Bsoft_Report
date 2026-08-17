from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from build_standard_knowledge_pack import build_pack  # noqa: E402
from knowledge_db import sha256_file, sha256_text  # noqa: E402
from postgres_knowledge_db import canonical_json  # noqa: E402
from run_project_pipeline import run_pipeline  # noqa: E402


def make_company_knowledge(root: Path) -> Path:
    solution = root / "company-standard.docx"
    document = Document()
    document.add_heading("建设内容", level=1)
    document.add_heading("电子病历系统", level=2)
    document.add_heading("临床文书与质控", level=3)
    document.add_paragraph(
        "电子病历系统面向门诊和住院临床业务，统一患者身份、文书模板、病历数据和质量控制规则。"
        "系统接收患者、就诊、医嘱和诊断等业务数据，形成文书创建、签名、归档、质控反馈和整改闭环。"
        "建设时应明确与HIS、集成平台、医技系统和电子签名服务的接口关系，并以流程测试、数据一致性检查、权限审计和质控记录作为验收证据。"
    )
    document.add_heading("病历质量控制", level=4)
    document.add_paragraph(
        "病历质量控制支持时限、完整性、逻辑一致性和缺陷整改规则，按科室、人员和问题类型形成可追溯记录。"
        "规则启用范围、提醒方式、申诉处理和统计口径应由医院管理部门确认。"
    )
    document.save(solution)

    scope = root / "company-scope.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "标准能力"
    sheet.append(["大类", "系统名称", "模块名称", "产品功能", "三级医院推荐"])
    sheet.append(["临床业务", "电子病历系统", "临床文书与质控", "病历质量控制", "推荐"])
    workbook.save(scope)
    payload = build_pack(solution, scope, title="公司医疗信息化标准知识包")
    pack_path = root / "reviewed-standard-pack.json"
    pack_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return pack_path


def make_project_scope(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "建设清单"
    sheet.append(["序号", "建设内容", "分类", "建设方式", "数量", "单位", "验收目标"])
    sheet.append([1, "电子病历系统", "临床应用", "升级", 1, "套", "形成临床文书与病历质控闭环"])
    workbook.save(path)


class LocalV1EndToEndTests(unittest.TestCase):
    def test_new_project_uses_company_pack_to_generate_construction_chapter_and_full_working_report(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pack_path = make_company_knowledge(root)
            project = root / "new-project"
            source_dir = project / "原始资料"
            source_dir.mkdir(parents=True)
            make_project_scope(source_dir / "项目建设清单.xlsx")
            (source_dir / "项目说明.md").write_text(
                "# 项目说明\n\n本项目拟升级电子病历系统，具体投资、工期和现状数据待院方进一步确认。\n",
                encoding="utf-8",
            )

            result = run_pipeline(
                project,
                project_code="LOCAL-V1-E2E",
                official_name="测试医院电子病历升级项目",
                owner_name="测试医院",
                knowledge_mode="offline_pack",
                standard_knowledge_packs=[pack_path],
            )

            manifest = result["standard_knowledge"]["manifest"]
            coverage = result["construction_knowledge_coverage"]
            self.assertEqual(manifest["source"], "reviewed_offline_pack")
            self.assertGreater(manifest["counts"]["corpus_blocks"], 0)
            self.assertGreater(manifest["counts"]["capabilities"], 0)
            self.assertEqual(manifest["connection_status"], "not_required")
            self.assertEqual(
                manifest["permission_scopes"][manifest["package_ids"][0]],
                "internal_company_reuse",
            )
            self.assertTrue(manifest["synced_at"][manifest["package_ids"][0]])
            imported = result["standard_knowledge"]["imports"][0]
            pack_payload = json.loads(pack_path.read_text(encoding="utf-8"))
            self.assertEqual(imported["content_hash"], sha256_text(canonical_json(pack_payload)))
            self.assertEqual(imported["file_sha256"], sha256_file(pack_path))
            self.assertEqual(
                manifest["content_hashes"][manifest["package_ids"][0]],
                imported["content_hash"],
            )
            self.assertEqual(coverage["scope_count"], 1)
            self.assertEqual(coverage["scope_node_count"], 1)
            self.assertEqual(coverage["mapped_scope_count"], 1)
            self.assertGreater(coverage["standard_block_count"], 0)
            self.assertGreater(result["draft_generation"]["total_visible_length"], 20000)
            self.assertEqual(
                result["draft_generation"]["working_draft_adopted_count"],
                result["draft_generation"]["adopted_count"],
            )
            self.assertEqual(result["draft_generation"]["formal_delivery_adopted_count"], 0)
            self.assertIn("working-report assembly only", result["draft_generation"]["adoption_notice"])

            report = (project / "11-正文工作稿" / "report-working.md").read_text(encoding="utf-8")
            self.assertIn("第5章 建设内容", report)
            self.assertIn("电子病历系统", report)
            self.assertIn("质控反馈和整改闭环", report)
            self.assertIn("建设内容总表", report)
            self.assertIn("【工作稿】章节“采纳”仅表示已选入本轮工作稿组装", report)

            construction_packages = sorted((project / "10-章节任务包").glob("CH5.1.1-*.json"))
            self.assertEqual(len(construction_packages), 1)
            package = json.loads(construction_packages[0].read_text(encoding="utf-8"))
            corpus_sources = [
                item["source_object_id"] for item in package["sources"] if item["source_type"] == "corpus"
            ]
            self.assertEqual(corpus_sources, coverage["block_ids"])
            self.assertTrue((project / "数据包" / "结构化数据" / "construction-knowledge-coverage.json").is_file())

            rerun = run_pipeline(
                project,
                project_code="LOCAL-V1-E2E",
                official_name="测试医院电子病历升级项目",
                owner_name="测试医院",
                knowledge_mode="offline_pack",
                standard_knowledge_packs=[pack_path],
            )
            rerun_manifest = rerun["standard_knowledge"]["manifest"]
            self.assertEqual(rerun_manifest["package_ids"], manifest["package_ids"])
            self.assertEqual(rerun_manifest["content_hashes"], manifest["content_hashes"])
            self.assertEqual(rerun_manifest["counts"], manifest["counts"])
            self.assertEqual(rerun_manifest["counts"]["policy_clauses"], 0)


if __name__ == "__main__":
    unittest.main()
