import sys
import tempfile
import unittest
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/"scripts"))
from lint_docx_format import lint_docx


class ConstructionLintTests(unittest.TestCase):
    def test_toc_tab_and_number_are_not_fake_headings_but_body_is_checked(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)/"toc.docx"
            doc=Document()
            doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
            paragraph=doc.add_paragraph("1.1 目录项\t3",style="TOC 1")
            paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(5))
            doc.add_paragraph("1.1 假标题\t3")
            doc.save(path)
            issues=lint_docx(path)["issues"]
            false_issues={"heading_like_without_heading_style","direct_tab_stops","leading_whitespace"}
            self.assertFalse(any(i["paragraph_index"]==1 and i["issue_type"] in false_issues for i in issues))
            self.assertTrue(any(i["paragraph_index"]==2 and i["issue_type"]=="heading_like_without_heading_style" for i in issues))


if __name__=="__main__":unittest.main()
