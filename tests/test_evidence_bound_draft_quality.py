from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from benchmark_report_quality import docx_content_profile, evaluate, template_skeleton  # noqa: E402
from build_dynamic_construction_outline import presentation_title  # noqa: E402
from build_evidence_bound_initial_drafts import (  # noqa: E402
    CHAPTER_ARGUMENT_OUTLINES,
    ensure_minimum_length,
    feature_control_profile,
    generic_section,
    policy_background_section,
    policy_section,
)
from validate_section_draft import assess_content, visible_length  # noqa: E402


class EvidenceBoundDraftQualityTests(unittest.TestCase):
    def test_missing_content_remains_a_validation_gap_instead_of_padding(self) -> None:
        lines = ensure_minimum_length(
            ["## 项目基本情况", "", "已有项目事实。", ""],
            {"official_name": "测试医院信息化建设项目"},
            {
                "chapter_code": "1.1.1",
                "section_title": "项目基本情况",
                "section_role": "project_overview",
                "length_min": 500,
            },
            [{"standard_name": "门诊业务系统"}, {"standard_name": "数据治理平台"}],
        )
        content = "\n".join(lines)
        self.assertLess(visible_length(content), 500)
        assessment = assess_content(content, {"chapter_code": "1.1.1", "length_min": 500}, [], mode="working")
        self.assertIn("section_too_short", {issue["code"] for issue in assessment["issues"]})

    def test_construction_content_is_never_padded_or_changed(self) -> None:
        original = ["## 业务应用系统建设", "", "<!-- corpus-block id=BLOCK-TEST -->", "经确认的标准正文。", ""]
        for minimum in (0, 100, 10000):
            plan = {"chapter_code": "5.1.1", "section_role": "construction_content", "length_min": minimum}
            self.assertEqual(ensure_minimum_length(original.copy(), {}, plan, []), original)

    def test_generic_chapter_has_project_specific_structure_without_template_residue(self) -> None:
        project = {"official_name": "测试医院信息化建设项目"}
        plan = {
            "chapter_code": "6.1.1",
            "section_title": "项目组织与进度安排",
            "section_role": "implementation_operation",
        }
        scopes = [
            {"standard_name": "电子病历系统", "construction_mode": "upgrade"},
            {"standard_name": "医院信息集成平台", "construction_mode": "new_build"},
        ]
        content = "\n".join(generic_section(project, plan, scopes, []))
        self.assertIn("项目治理与组织分工", content)
        self.assertIn("阶段计划与成果物", content)
        for marker in (
            "项目数据库",
            "公司能力",
            "标准方案可复用的机制包括",
            "管理与实施机制",
            "当前登记的",
            "本节结合",
            "可评审初稿",
            "工作稿按",
            "候选功能",
        ):
            self.assertNotIn(marker, content)

    def test_major_chapters_use_distinct_argument_outlines(self) -> None:
        self.assertNotEqual(
            CHAPTER_ARGUMENT_OUTLINES["6.1.1"],
            CHAPTER_ARGUMENT_OUTLINES["6.1.2"],
        )
        self.assertNotEqual(
            CHAPTER_ARGUMENT_OUTLINES["7.1.2"],
            CHAPTER_ARGUMENT_OUTLINES["8.1.1"],
        )

    def test_feature_profiles_prioritize_business_semantics(self) -> None:
        nursing_design, _ = feature_control_profile("护士临床路径", "医嘱执行与交接")
        reminder_design, _ = feature_control_profile("业务提醒", "接口消息触发提醒")
        self.assertIn("患者身份", nursing_design)
        self.assertNotIn("开立、审核", nursing_design)
        self.assertIn("规则来源", reminder_design)
        self.assertNotIn("数据提供方", reminder_design)

    def test_four_basis_groups_and_three_level_policy_background_are_generated(self) -> None:
        project = {"official_name": "测试医院信息化建设项目"}
        material = {
            "match_run_id": "POLICYMATCHRUN-TEST",
            "material_signature": "a" * 64,
            "basis_groups": {
                "policy_basis": [{
                    "policy_id": "POLICY-1",
                    "clause_ids": ["CLAUSE-1"],
                    "title": "公立医院高质量发展意见",
                    "document_no": "国办发〔2021〕18号",
                }],
                "industry_standard": [{
                    "policy_id": "STANDARD-1",
                    "clause_ids": ["CLAUSE-2"],
                    "title": "电子病历系统应用水平分级评价标准",
                    "document_no": "国卫办医函〔2018〕1079号",
                }],
                "security_standard": [{
                    "policy_id": "SECURITY-1",
                    "clause_ids": ["CLAUSE-3"],
                    "title": "信息安全技术 网络安全等级保护基本要求",
                    "document_no": "GB/T 22239—2019",
                }],
                "investment_basis": [{
                    "policy_id": "INVESTMENT-1",
                    "clause_ids": ["CLAUSE-4"],
                    "title": "软件工程 软件开发成本度量规范",
                    "document_no": "GB/T 36964—2018",
                }],
            },
            "working_basis_groups": {
                "policy_basis": [{
                    "title": "《十四五全民健康信息化规划》",
                    "document_no": "国卫规划发〔2022〕30号",
                }],
                "industry_standard": [],
                "security_standard": [],
                "investment_basis": [],
            },
            "policy_background_groups": {
                "national": [{
                    "policy_id": "POLICY-1",
                    "clause_ids": ["CLAUSE-1"],
                    "text_hash": "b" * 64,
                    "text": "《公立医院高质量发展意见》提出推进智慧医院和医院信息标准化建设。",
                }],
                "province": [],
                "prefecture": [],
            },
            "policy_background_candidate_groups": {
                "national": [{"title": "《十四五全民健康信息化规划》"}],
                "province": [],
                "prefecture": [],
            },
            "working_policy_background_is_ordered_subsequence": True,
            "quality": {
                "formal_basis": {
                    section: {"gap": 0}
                    for section in (
                        "policy_basis", "industry_standard", "security_standard", "investment_basis"
                    )
                },
                "formal_background": {
                    "national": {"gap": 0},
                    "province": {"gap": 2},
                    "prefecture": {"gap": 1},
                },
                "candidate_basis": {"groups": {}},
                "delivery_blockers": ["formal_policy_background_province_gap:2"],
            },
        }
        basis = "\n".join(policy_section(
            project,
            {"chapter_code": "1.2.1", "section_title": "可行性研究报告编制依据"},
            material,
            ["编制依据表"],
        ))
        governance = "\n".join(policy_section(
            project,
            {"chapter_code": "1.2.2", "section_title": "编制依据适用与动态更新"},
            material,
            [],
        ))
        background = "\n".join(policy_background_section(
            project,
            {"chapter_code": "2.1.1", "section_title": "政策背景"},
            material,
        ))
        for heading in ("政策类依据", "行业标准依据", "安全类标准依据", "投资估算编制依据"):
            self.assertIn(heading, basis)
        self.assertIn("公立医院高质量发展意见", basis)
        self.assertIn("电子病历系统应用水平分级评价标准", basis)
        self.assertIn("网络安全等级保护基本要求", basis)
        self.assertIn("软件开发成本度量规范", basis)
        self.assertIn("十四五全民健康信息化规划", basis)
        self.assertIn("【待核验】", basis)
        self.assertIn("| 序号 | 依据名称 | 文号/标准号 | 使用状态 |", basis)
        self.assertIn('match-run="POLICYMATCHRUN-TEST"', basis)
        self.assertIn('signature="' + "a" * 64 + '"', basis)
        self.assertIn('policy-item id="POLICY-1" clauses="CLAUSE-1"', basis)
        basis_assessment = assess_content(
            basis,
            {
                "chapter_code": "1.2.1",
                "section_role": "basis",
                "required_tables_json": '["编制依据表"]',
            },
            [],
            mode="working",
        )
        self.assertNotIn(
            "required_table_missing",
            {issue["code"] for issue in basis_assessment["issues"]},
        )
        self.assertIn("#### 1.2.2.4 版本更新机制", governance)
        self.assertNotIn("电子病历系统应用水平分级评价标准", governance)
        for heading in ("国家政策背景", "省/自治区政策背景", "市/项目建设地区政策背景"):
            self.assertIn(heading, background)
        self.assertIn("提出推进智慧医院", background)
        self.assertIn('policy-background id="POLICY-1" clauses="CLAUSE-1"', background)
        self.assertIn("仅列为核验任务", background)
        self.assertNotIn("电子病历系统应用水平分级评价标准", background)
        self.assertNotIn("引用《", basis + governance + background)
        self.assertNotIn("正式报审前还需复核", basis + governance + background)

    def test_public_outline_titles_remove_company_branding(self) -> None:
        title = presentation_title("创业慧康 BsoftGPT 医院信息平台")
        self.assertNotIn("创业慧康", title)
        self.assertNotIn("Bsoft", title)
        self.assertIn("大模型智能", title)

    def test_benchmark_blocks_template_prose(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "candidate.docx"
            document = Document()
            document.add_heading("第1章 总论", level=1)
            document.add_paragraph("项目数据库中的范围记录用于形成正文。" * 8)
            document.add_paragraph("本项目围绕临床业务、数据治理和运行保障形成完整论证。" * 8)
            document.add_paragraph("各项建设内容均应建立设计、测试、验收和持续改进证据。" * 8)
            document.save(path)
            result = evaluate(path)
        self.assertIn(
            "template_prose_residue",
            {item["code"] for item in result["blockers"]},
        )

    def test_benchmark_uses_four_basis_groups_and_confirmed_hard_minimums(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "four-groups.docx"
            document = Document()
            for heading, title in (
                ("政策类依据", "测试政策文件"),
                ("行业标准依据", "测试行业标准"),
                ("安全类标准依据", "测试安全标准"),
                ("投资估算编制依据", "测试投资依据"),
            ):
                document.add_heading(heading, level=4)
                document.add_paragraph(f"《{title}》")
            document.save(path)
            profile = docx_content_profile(path)
            result = evaluate(path)

        self.assertEqual(
            profile["basis_group_title_mentions"],
            {
                "policy_basis": 1,
                "industry_standard": 1,
                "security_standard": 1,
                "investment_basis": 1,
            },
        )
        codes = {item["code"] for item in result["blockers"]}
        self.assertTrue(
            {
                "policy_basis_hard_minimum",
                "industry_standard_hard_minimum",
                "security_standard_hard_minimum",
                "investment_basis_hard_minimum",
            }.issubset(codes)
        )

    def test_template_skeleton_normalizes_topic_and_scope_substitution(self) -> None:
        first = (
            "针对项目组织与进度安排，实施对象包括电子病历、集成平台、数据中心、患者服务平台等20项建设内容，"
            "计划编排时需识别公共依赖。从建设方式看，升级事项包括电子病历、集成平台、数据中心、患者服务平台。"
        )
        second = (
            "针对投资分项估算，实施对象包括PACS、决策支持、数据治理、移动护理等16项建设内容，"
            "计划编排时需识别公共依赖。从建设方式看，新建事项包括PACS、决策支持、数据治理、移动护理。"
        )
        self.assertEqual(template_skeleton(first), template_skeleton(second))


if __name__ == "__main__":
    unittest.main()
