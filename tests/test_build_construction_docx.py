from __future__ import annotations

import copy
import sys
import unittest
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import test_construction_alignment as fixture
from construction_alignment import assemble
from build_construction_docx import audit_heading_paths, build_construction_docx


class ConstructionWordTests(unittest.TestCase):
    def setUp(self):
        self.fixture = fixture.ConstructionAlignmentTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.tearDown)
        matched, reviewed = self.fixture._flattened_review_case()
        self.matched, self.reviewed = matched, reviewed
        self.manifest, _ = assemble(self.fixture.database, "ALIGN-001", matched["match_run_id"],
                                    hierarchy_review=reviewed)
        self.output = self.fixture.database.parent / "construction.docx"

    def test_bound_word_uses_builtin_preset_and_actual_parent_styles(self):
        result = build_construction_docx(self.fixture.database, self.manifest, self.output, "合成层级测试")
        self.assertTrue(result["structural_validation_passed"])
        self.assertTrue(result["heading_ancestry_audit"]["passed"])
        self.assertEqual(result["visual_render_review"], "not_run")
        doc = Document(self.output)
        headings = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading ")]
        self.assertEqual([text for style, text in headings if style == "Heading 1"], ["建设清单", "建设内容"])
        self.assertIn(("Heading 2", "院内集成平台及数据中心"), headings)
        self.assertIn(("Heading 3", "医院信息基础平台"), headings)
        self.assertIn(("Heading 4", "患者主索引系统"), headings)
        self.assertNotIn(("Heading 2", "患者主索引系统"), headings)
        self.assertEqual(doc.core_properties.title, "合成层级测试建设清单与建设内容")
        self.assertEqual(sum(p.text == "【待补充】" for p in doc.paragraphs), 1)
        self.assertIn("患者主索引系统", [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
        self.assertEqual(doc.styles["Heading 2"].element.xpath("./w:rPr/w:rFonts")[0].get(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia"), "黑体")

    def test_word_parent_promotion_is_detected_independently(self):
        build_construction_docx(self.fixture.database, self.manifest, self.output, "合成层级测试")
        doc = Document(self.output)
        next(p for p in doc.paragraphs if p.text == "患者主索引系统").style = "Heading 2"
        doc.save(self.output)
        with self.assertRaisesRegex(ValueError, "ancestry|ancestor"):
            audit_heading_paths(self.output, self.manifest)

    def test_stale_manifest_cannot_produce_word(self):
        tampered = copy.deepcopy(self.manifest)
        tampered["application_software_solution"]["items"][0]["display_parent_path"] = ["变造目录"]
        with self.assertRaisesRegex(ValueError, "validation"):
            build_construction_docx(self.fixture.database, tampered, self.output, "合成层级测试")
        self.assertFalse(self.output.exists())

    def test_word_navigation_outline_is_checked_not_just_style_name(self):
        build_construction_docx(self.fixture.database, self.manifest, self.output, "合成层级测试")
        doc = Document(self.output)
        from docx.oxml.ns import qn
        doc.styles["Heading 4"].element.find(qn("w:pPr")).find(qn("w:outlineLvl")).set(qn("w:val"), "1")
        doc.save(self.output)
        with self.assertRaisesRegex(ValueError, "outline level"):
            audit_heading_paths(self.output, self.manifest)

    def test_overdeep_standard_subtree_is_not_flattened(self):
        reviewed = copy.deepcopy(self.reviewed)
        for item in reviewed["items"]:
            item["parent_path"] = ["一级", "二级", "三级", "四级"]
        with self.assertRaisesRegex(ValueError, "heading_depth_exceeded"):
            assemble(self.fixture.database, "ALIGN-001", self.matched["match_run_id"], hierarchy_review=reviewed)


if __name__ == "__main__":
    unittest.main()
