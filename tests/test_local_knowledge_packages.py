from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from build_standard_knowledge_pack import build_pack  # noqa: E402
from init_project_workbench import initialize_project  # noqa: E402
from knowledge_db import sha256_file, sha256_text  # noqa: E402
from local_knowledge_packages import (  # noqa: E402
    LocalKnowledgeError,
    build_policy_package,
    build_standard_package,
    install_local_package,
    load_config,
    package_status,
    query_packages,
    sync_to_project,
)


def write_json(path: Path, value: dict) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def make_standard_pack(root: Path, *, suffix: str = "") -> Path:
    solution = root / f"standard{suffix}.docx"
    document = Document()
    document.add_heading("建设内容", level=1)
    document.add_heading(f"电子病历系统{suffix}", level=2)
    document.add_paragraph(
        "电子病历系统围绕临床文书、病历质控、数据共享和闭环管理形成业务能力。" * 8
    )
    document.save(solution)

    scope = root / f"scope{suffix}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["大类", "系统名称", "模块名称", "产品功能"])
    sheet.append(["临床业务", "电子病历系统", f"电子病历系统{suffix}", "病历质控"])
    workbook.save(scope)

    payload = build_pack(solution, scope, title=f"智慧医院标准方案{suffix}")
    output = root / f"standard-pack{suffix}.json"
    write_json(output, payload)
    return output


def policy_catalog() -> dict:
    row = {
        "catalog_entry_id": "POLICYCATENTRY-LOCAL-1",
        "source_row": 1,
        "source_index_no": "CN1.001",
        "identity_key": "POLICYIDENTITY-LOCAL-1",
        "catalog_group_code": "CN1",
        "catalog_group_name": "国家政策",
        "authority_level_label": "国家",
        "category_name": "信息化",
        "keyword_text": "智慧医院 电子病历",
        "keyword_tags": ["智慧医院", "电子病历"],
        "document_no": "国办发〔2021〕18号",
        "title": "《关于推动公立医院高质量发展的意见》",
        "publish_date": "2021-06-04",
        "publish_date_raw": "2021-06-04",
        "issuer": "国务院办公厅",
        "file_count": 1,
        "notes": "",
        "external_url": "https://www.gov.cn/zhengce/content/2021-06/04/content_5615473.htm",
        "verification_status": "unverified",
        "entry_status": "active",
    }
    row["row_hash"] = sha256_text(json.dumps(row, ensure_ascii=False, sort_keys=True))
    return {
        "schema_version": "1.0",
        "catalog_id": "POLICYCATALOG-LOCAL-PACKAGE-TEST",
        "catalog_scope": "medical_health_national",
        "title": "医疗卫生政策目录测试",
        "permission_scope": "internal_company_reference",
        "source_file": {"file_name": "policy.xlsx", "sha256": "a" * 64},
        "worksheet_name": "政策",
        "records": [row],
        "built_at": "2026-09-02T00:00:00+08:00",
        "content_hash": "b" * 64,
    }


def document_standards() -> dict:
    return {
        "seed_version": "test-v1",
        "standards": [
            {
                "standard_id": "DOCSTD-LOCAL-TEST",
                "document_type": "feasibility_study",
                "jurisdiction_code": "100000",
                "jurisdiction_name": "全国",
                "authority_name": "国家发展改革委",
                "title": "政府投资项目可行性研究报告编写通用大纲（测试）",
                "version": "2023",
                "effective_date": "2023-05-01",
                "expiry_date": "",
                "required_sections": ["项目概述", "项目建设背景和必要性"],
                "optional_sections": [],
                "table_requirements": ["投资估算及资金筹措"],
                "official_url": "https://example.gov.cn/standard",
                "verification_status": "verified",
                "status": "active",
            }
        ],
    }


def make_config(root: Path, standard: Path, policy: Path) -> Path:
    path = root / "medical-report-local-kb.json"
    write_json(
        path,
        {
            "schema_version": "1.0",
            "packages": {
                "standard": {"path": str(standard), "max_age_days": 45, "required": True},
                "policy": {"path": str(policy), "max_age_days": 8, "required": True},
            },
        },
    )
    return path


class LocalKnowledgePackageTests(unittest.TestCase):
    def test_build_query_and_sync_split_packages(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_pack = make_standard_pack(root)
            catalog = root / "policy-catalog.json"
            standards = root / "document-standards.json"
            write_json(catalog, policy_catalog())
            write_json(standards, document_standards())
            standard_db = root / "standard-knowledge.sqlite"
            policy_db = root / "policy-knowledge.sqlite"

            standard_result = build_standard_package(
                source_pack, standard_db, release_version="2026.09"
            )
            policy_result = build_policy_package(
                policy_db,
                release_version="2026-W36",
                catalog_path=catalog,
                document_standards_path=standards,
            )
            self.assertEqual(standard_result["release"]["knowledge_kind"], "standard")
            self.assertEqual(policy_result["release"]["knowledge_kind"], "policy")
            self.assertEqual(policy_result["warnings"], ["formal_policy_clauses_empty"])

            config = load_config(make_config(root, standard_db, policy_db))
            status = package_status(config)
            self.assertEqual(status["status"], "ready")
            self.assertTrue(status["remote_mcp_retained"])
            corpus = query_packages(config, kind="corpus", search="电子病历", limit=10)
            catalog_rows = query_packages(
                config, kind="policy-catalog", search="高质量发展", limit=10
            )
            standard_rows = query_packages(
                config,
                kind="document-standard",
                document_type="feasibility_study",
                limit=10,
            )
            self.assertGreater(corpus["count"], 0)
            self.assertEqual(catalog_rows["count"], 1)
            self.assertEqual(standard_rows["count"], 1)

            initialized = initialize_project(root / "project", project_code="LOCAL-KB-001")
            synced = sync_to_project(
                config,
                project_database=Path(initialized["database"]),
                project_code="LOCAL-KB-001",
            )
            counts = synced["snapshot_validation"]["counts"]
            self.assertGreater(counts["corpus_blocks"], 0)
            self.assertGreater(counts["capabilities"], 0)
            self.assertEqual(counts["catalog_records"], 1)
            self.assertEqual(counts["document_standards"], 1)

            newer_source = make_standard_pack(root, suffix="新版")
            newer_db = root / "standard-new.sqlite"
            build_standard_package(newer_source, newer_db, release_version="2026.10")
            newer_config = load_config(make_config(root, newer_db, policy_db))
            with self.assertRaisesRegex(LocalKnowledgeError, "pinned"):
                sync_to_project(
                    newer_config,
                    project_database=Path(initialized["database"]),
                    project_code="LOCAL-KB-001",
                )

    def test_install_requires_digest_and_keeps_previous_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            first = root / "first.sqlite"
            second = root / "second.sqlite"
            build_standard_package(make_standard_pack(root), first, release_version="2026.09")
            build_standard_package(
                make_standard_pack(root, suffix="升级"), second, release_version="2026.10"
            )
            installed = root / "installed" / "standard-knowledge.sqlite"
            placeholder_policy = root / "policy-knowledge.sqlite"
            config = load_config(make_config(root, installed, placeholder_policy))

            first_result = install_local_package(
                config, kind="standard", candidate=first, expected_sha256=sha256_file(first)
            )
            second_result = install_local_package(
                config, kind="standard", candidate=second, expected_sha256=sha256_file(second)
            )
            self.assertEqual(first_result["status"], "installed")
            self.assertEqual(second_result["status"], "installed")
            self.assertTrue(Path(second_result["backup"]).is_file())
            self.assertEqual(sha256_file(installed), sha256_file(second))
            with self.assertRaisesRegex(LocalKnowledgeError, "does not match"):
                install_local_package(
                    config, kind="standard", candidate=first, expected_sha256="0" * 64
                )


if __name__ == "__main__":
    unittest.main()
