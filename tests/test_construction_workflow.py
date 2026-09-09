from __future__ import annotations
import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from test_construction_alignment import ConstructionAlignmentTests
from test_local_knowledge_packages import make_standard_pack
from construction_alignment import capture_scope_snapshot, match_scope, apply_decisions, assemble, validate_manifest
from construction_model import derive_scope, traceability, build_heading_tree
from construction_workflow import prepare_review, confirm_and_generate, resume, get_status, read, write
from construction_word import audit_word
from local_knowledge_packages import build_standard_package, load_config, sync_to_project
from knowledge_db import connect, apply_migrations, upsert_project


class ConstructionOptimizationTests(unittest.TestCase):
    def setUp(self):
        self.f = ConstructionAlignmentTests(); self.f.setUp()
        self.addCleanup(self.f.tearDown)

    def test_traceability_expansion_heading_and_duplicate_preserve_raw_rows(self):
        original = copy.deepcopy(self.f.scope_payload)
        plan = [{"original_ordinal": 1, "role": "heading_only", "title": "平台分组", "parent_path": []},
                {"original_ordinal": 2, "parent_path": ["平台分组"], "modules": [{"name": "主数据管理"}, {"name": "主数据管理"}]}]
        payload = derive_scope(original, plan)
        from ingest_scope_items_sqlite import ingest_scope_payload
        ingest_scope_payload(self.f.database, payload, project_code="ALIGN-001")
        snap = capture_scope_snapshot(self.f.database, "ALIGN-001", payload)
        proof = traceability(payload, snap["scope_rows"])
        self.assertEqual(proof["original_row_count"], 3)
        self.assertEqual(proof["assembly_item_count"], 3)
        self.assertEqual(proof["heading_only_count"], 1)
        self.assertEqual([e["display_ref"] for e in proof["mapping_edges"]], ["2.1", "2.2", "3"])
        matched = match_scope(self.f.database, "ALIGN-001")
        self.assertEqual(len(matched["duplicate_mappings"]), 1)
        self.assertEqual(len(matched["duplicate_mappings"][0]["items"]), 2)
        decisions = [{"scope_row_id": i["scope_row_id"], "decision": "confirmed", "candidate_id": i["candidates"][0]["candidate_id"]} if i["candidates"] else {"scope_row_id": i["scope_row_id"], "decision": "confirmed_gap"} for i in matched["items"]]
        apply_decisions(self.f.database, {"match_run_id": matched["match_run_id"], "reviewed_by": "tester", "decisions": decisions})
        manifest, text = assemble(self.f.database, "ALIGN-001", matched["match_run_id"])
        self.assertTrue(validate_manifest(self.f.database, manifest)["valid"])
        self.assertEqual(manifest["construction_list_import"]["original_display_payload"], original)
        self.assertEqual(text.count("主数据管理标准总述，必须原文导入。"), 2)
        self.assertIn("平台分组", [n.get("title") for n in manifest["heading_tree"]["nodes"]])
        tampered = copy.deepcopy(manifest)
        tampered["application_software_solution"]["items"].pop()
        self.assertIn("scope_row_coverage_mismatch", [i["code"] for i in validate_manifest(self.f.database, tampered)["issues"]])
        altered = copy.deepcopy(payload); altered["original_rows"][0]["assembly_count"] = 1
        with self.assertRaises(ValueError): traceability(altered, snap["scope_rows"])

    def test_match_reuse_and_decision_retry_are_idempotent(self):
        capture_scope_snapshot(self.f.database, "ALIGN-001", self.f.scope_payload)
        first = match_scope(self.f.database, "ALIGN-001")
        second = match_scope(self.f.database, "ALIGN-001")
        self.assertEqual(first, second)
        payload = {"match_run_id": first["match_run_id"], "reviewed_by": "tester", "decisions": [{"scope_row_id": first["items"][2]["scope_row_id"], "decision": "confirmed_gap"}]}
        a = apply_decisions(self.f.database, payload)
        b = apply_decisions(self.f.database, {**payload, "reviewed_at": "2026-09-10T10:00:00+08:00"})
        self.assertEqual((a["applied"], b["applied"], b["duplicates"]), (1, 0, 1))
        changed = match_scope(self.f.database, "ALIGN-001", similar_threshold=.7)
        self.assertNotEqual(first["match_run_id"], changed["match_run_id"])

    def test_latest_decisions_invalidate_old_manifest(self):
        matched, hierarchy = self.f._flattened_review_case()
        manifest, _ = assemble(self.f.database, "ALIGN-001", matched["match_run_id"], hierarchy_review=hierarchy)
        apply_decisions(self.f.database, {"match_run_id": matched["match_run_id"], "reviewed_by": "tester", "decisions": [{"scope_row_id": matched["items"][0]["scope_row_id"], "decision": "confirmed_gap"}]})
        self.assertFalse(validate_manifest(self.f.database, manifest)["valid"])

    def test_77_original_rows_expand_to_80_without_losing_heading_or_gaps(self):
        from ingest_scope_items_sqlite import ingest_scope_payload
        original = copy.deepcopy(self.f.scope_payload)
        sheet = original["sheets"][0]
        sheet["rows"] = [{"_source_row": n+1, "序号": str(n), "软件大类": "合成业务分组", "软件系统名称": "医院信息基础平台", "模块名称": "合成缺口" if n in {58,72,73,74,75,76,77} else "主数据管理"} for n in range(1,78)]
        plan = [{"original_ordinal": 3, "modules": [{"name": "主数据管理"} for _ in range(5)]},
                {"original_ordinal": 37, "role": "heading_only", "title": "合成标题组", "parent_path": []}]
        derived = derive_scope(original, plan)
        ingest_scope_payload(self.f.database, derived, project_code="ALIGN-001")
        snap = capture_scope_snapshot(self.f.database, "ALIGN-001", derived)
        matched = match_scope(self.f.database, "ALIGN-001")
        decisions = [{"scope_row_id": i["scope_row_id"], "decision": "confirmed", "candidate_id": i["candidates"][0]["candidate_id"]} if i["candidates"] else {"scope_row_id": i["scope_row_id"], "decision": "confirmed_gap"} for i in matched["items"]]
        apply_decisions(self.f.database, {"match_run_id": matched["match_run_id"], "reviewed_by": "test-user", "decisions": decisions})
        manifest, _ = assemble(self.f.database, "ALIGN-001", matched["match_run_id"])
        self.assertTrue(validate_manifest(self.f.database, manifest)["valid"])
        self.assertEqual((manifest["traceability"]["original_row_count"], manifest["traceability"]["assembly_item_count"]), (77,80))
        self.assertEqual(sum(i["status"] == "verbatim" for i in manifest["application_software_solution"]["items"]),73)
        self.assertEqual(sum(i["status"] == "pending_supplement" for i in manifest["application_software_solution"]["items"]),7)
        self.assertEqual(len(manifest["construction_list_import"]["original_display_payload"]["sheets"][0]["rows"]),77)


class ConstructionWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        pack = make_standard_pack(self.root)
        standard = self.root / "standard-knowledge.sqlite"
        build_standard_package(pack, standard, release_version="test-v1")
        self.config = self.root / "local.json"
        write(self.config, {"schema_version": "1.0", "packages": {"standard": {"path": str(standard)}}})
        self.scope = self.root / "input.csv"
        self.scope.write_text("序号,大类,系统名称,模块名称\n1,临床业务,电子病历系统,电子病历系统\n2,临床业务,未知系统,未知模块\n", encoding="utf-8")
        self.work = self.root / "work"
        self.plan = [{"original_ordinal": 1, "parent_path": ["临床业务"]}, {"original_ordinal": 2, "parent_path": ["临床业务"]}]

    def prepare(self):
        return prepare_review(self.work, self.scope, "TEST", plan=self.plan, local_config=self.config, render_after_build=False)

    def confirm(self, status):
        return confirm_and_generate(self.work, {"review_id": status["review_id"], "reviewed_by": "test-user", "user_reply": "全部按建议处理并生成Word", "accept_all": True})

    def test_standard_only_prepare_confirm_docx_and_resume(self):
        status = self.prepare()
        self.assertEqual(status["status"], "awaiting_confirmation")
        self.assertEqual(list(self.work.glob("*.md")), [self.work / "01-请确认建设清单对照.md"])
        self.assertFalse(list(self.work.rglob("*.docx")))
        state = read(self.work / "运行数据/construction-state.json")
        with connect(Path(state["database"])) as conn:
            self.assertEqual(conn.execute("SELECT count(*) FROM policy_catalog_entry").fetchone()[0], 0)
        ready = self.confirm(status)
        self.assertEqual(ready["status"], "docx_created_render_pending")
        self.assertFalse(ready["delivery_ready"])
        docx = Path(ready["docx_path"])
        self.assertTrue(docx.is_file())
        before = docx.stat().st_mtime_ns
        with patch("construction_workflow.generate_word", side_effect=AssertionError("unnecessary Word rerun")), patch("construction_workflow.assemble", side_effect=AssertionError("unnecessary assembly rerun")):
            self.confirm(status); resume(self.work)
        self.assertEqual(before, docx.stat().st_mtime_ns)
        self.assertEqual(self.prepare()["review_id"], status["review_id"])
        manifest = read(self.work / "运行数据/construction-assembly-manifest.json")
        self.assertTrue(audit_word(docx, manifest)["valid"])
        document = Document(docx)
        from docx.oxml.ns import qn
        for p in document.paragraphs:
            if p.style.name.startswith("Heading "):
                props = p._p.pPr
                numbering = props.find(qn("w:numPr"))
                if numbering is not None:
                    props.remove(numbering)
        document.save(docx)
        self.assertTrue(audit_word(docx, manifest)["valid"], "Word may inherit numbering from heading styles")
        self.assertFalse(any("工作稿（未通过正式交付门禁）" in p.text for p in document.paragraphs))
        paragraph = next(p for p in document.paragraphs if "围绕临床文书" in p.text)
        paragraph.text = paragraph.text + "凭空增加内容"
        document.save(docx)
        self.assertFalse(audit_word(docx, manifest)["valid"])
        resume(self.work)
        self.assertTrue(audit_word(docx, manifest)["valid"])

    def test_stale_input_and_review_and_wrong_confirmation_block(self):
        status = self.prepare()
        with self.assertRaises(ValueError):
            self.confirm({"review_id": "old-review"})
        self.scope.write_text(self.scope.read_text(encoding="utf-8") + "3,新增,新增,新增\n", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "input changed"):
            self.confirm(status)

    def test_word_failure_resumes_without_reapplying_decisions(self):
        status = self.prepare()
        with patch("construction_workflow.generate_word", side_effect=OSError("renderer unavailable")):
            with self.assertRaises(OSError): self.confirm(status)
        self.assertEqual(get_status(self.work)["status"], "blocked")
        with patch("construction_workflow.apply_decisions", side_effect=AssertionError("must resume saved decisions")):
            result = resume(self.work)
        self.assertTrue(Path(result["docx_path"]).is_file())


if __name__ == "__main__": unittest.main()
