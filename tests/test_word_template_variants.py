from __future__ import annotations

import base64
import json
import sys
import tempfile
import unittest
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.shared import Cm


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from build_report_docx import build_docx  # noqa: E402
from knowledge_db import sha256_file  # noqa: E402


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Zl1sAAAAASUVORK5CYII="
)


class WordTemplateVariantTests(unittest.TestCase):
    def make_template(self, path: Path, font: str, margin_cm: float, marker: str) -> None:
        document = Document()
        normal = document.styles["Normal"]
        normal.font.name = "Arial"
        normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
        if "Body Text First Indent" not in {style.name for style in document.styles}:
            document.styles.add_style("Body Text First Indent", WD_STYLE_TYPE.PARAGRAPH)
        section = document.sections[0]
        section.left_margin = Cm(margin_cm)
        section.right_margin = Cm(margin_cm)
        document.add_paragraph(f"OLD-BODY-{marker}")
        document.add_picture(BytesIO(PNG_1X1))
        section.header.paragraphs[0].text = f"OLD-HEADER-{marker}"
        section.footer.paragraphs[0].text = f"OLD-FOOTER-{marker}"
        document.core_properties.author = f"OLD-AUTHOR-{marker}"
        document.save(path)

    def test_three_template_variants_preserve_styles_and_scrub_content(self) -> None:
        variants = [
            ("宋体", 2.0, "A"),
            ("仿宋", 2.5, "B"),
            ("楷体", 3.0, "C"),
        ]
        markdown = "# 模板测试项目\n\n## 可行性研究报告\n\n# 第1章 总论\n\n## 1.1 项目概况\n\n### 1.1.1 项目基本情况\n\n模板回归正文。\n"
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for font, margin_cm, marker in variants:
                template = root / f"template-{marker}.docx"
                output = root / f"output-{marker}.docx"
                self.make_template(template, font, margin_cm, marker)

                result = build_docx(
                    markdown,
                    output,
                    "模板测试项目",
                    "测试单位",
                    template=template,
                )

                self.assertEqual(result["template"], str(template.resolve()))
                generated = Document(output)
                self.assertAlmostEqual(generated.sections[0].left_margin.cm, margin_cm, places=1)
                east_asia = generated.styles["Normal"].element.rPr.rFonts.get(
                    qn("w:eastAsia")
                )
                self.assertEqual(east_asia, font)
                self.assertEqual(generated.core_properties.author, "")
                with zipfile.ZipFile(output) as package:
                    xml_text = "\n".join(
                        package.read(name).decode("utf-8", errors="ignore")
                        for name in package.namelist()
                        if name.endswith((".xml", ".rels"))
                    )
                    self.assertNotIn(f"OLD-BODY-{marker}", xml_text)
                    self.assertNotIn(f"OLD-HEADER-{marker}", xml_text)
                    self.assertNotIn(f"OLD-FOOTER-{marker}", xml_text)
                    self.assertNotIn(f"OLD-AUTHOR-{marker}", xml_text)
                    self.assertFalse(
                        any(name.startswith("word/media/") for name in package.namelist())
                    )

    def test_confirmed_format_config_controls_template_and_semantic_styles(self) -> None:
        markdown = "# 项目\n\n# 第1章 总论\n\n## 1.1 项目概况\n\n### 1.1.1 子标题\n\n模板正文。\n"
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            output = root / "output.docx"
            config = root / "word-format-authority.json"
            self.make_template(template, "仿宋", 2.5, "AUTHORITY")
            config.write_text(
                json.dumps(
                    {
                        "schema_version": "1.0",
                        "status": "confirmed",
                        "profile_id": "PROFILE-TEST",
                        "confirmed_by": "tester",
                        "confirmed_at": "2026-08-24T12:00:00+08:00",
                        "authority": {
                            "template_path": template.name,
                            "template_sha256": sha256_file(template),
                        },
                        "semantic_styles": {
                            "heading_1": "Heading 1",
                            "heading_2": "Heading 2",
                            "heading_3": "Heading 3",
                            "heading_4": "Heading 4",
                            "heading_5": "Heading 5",
                            "heading_6": "Heading 6",
                            "heading_7": "Heading 7",
                            "body": "Body Text First Indent",
                            "table": "Table Grid",
                            "table_header": "Normal",
                            "table_body": "Normal",
                        },
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result = build_docx(
                markdown,
                output,
                "模板测试项目",
                format_config=config,
            )

            self.assertEqual(result["format_profile_id"], "PROFILE-TEST")
            self.assertEqual(result["template_sha256"], sha256_file(template))
            generated = Document(output)
            body = next(paragraph for paragraph in generated.paragraphs if paragraph.text == "模板正文。")
            subtitle = next(paragraph for paragraph in generated.paragraphs if paragraph.text == "1.1.1 子标题")
            self.assertEqual(body.style.name, "Body Text First Indent")
            self.assertEqual(subtitle.style.name, "Heading 3")

    def test_default_preset_creates_and_explicitly_binds_multilevel_headings(self) -> None:
        markdown = (
            "# 项目\n\n"
            "# 第1章 总论\n\n"
            "## 1.1 项目概况\n\n"
            "### 1.1.1 项目基本情况\n\n"
            "#### 1.1.1.1 建设范围\n\n"
            "正文。\n"
        )
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "default-numbering.docx"
            result = build_docx(markdown, output, "默认编号测试项目")

            self.assertEqual(result["numbered_heading_levels"], list(range(1, 8)))
            self.assertEqual(result["numbered_heading_paragraphs"], 4)
            self.assertEqual(
                result["heading_numbering_audit"],
                {"status": "pass", "expected": 4, "explicitly_bound": 4},
            )

            document = Document(output)
            self.assertFalse(document.styles["Heading 4"].font.italic)
            expected = {
                "总论": 0,
                "项目概况": 1,
                "项目基本情况": 2,
                "建设范围": 3,
            }
            num_ids = set()
            for paragraph in document.paragraphs:
                if paragraph.text not in expected:
                    continue
                num_pr = paragraph._p.pPr.numPr
                self.assertIsNotNone(num_pr)
                self.assertEqual(num_pr.ilvl.val, expected[paragraph.text])
                self.assertGreater(num_pr.numId.val, 0)
                num_ids.add(num_pr.numId.val)
            self.assertEqual(len(num_ids), 1)
            with zipfile.ZipFile(output) as package:
                numbering_xml = package.read("word/numbering.xml").decode(
                    "utf-8", errors="ignore"
                )
            self.assertEqual(numbering_xml.count("<w:isLgl"), 6)

    def test_format_config_rejects_changed_profile_evidence(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            evidence = root / "format-profile.json"
            config = root / "word-format-authority.json"
            self.make_template(template, "仿宋", 2.5, "EVIDENCE")
            evidence.write_text("{}\n", encoding="utf-8")
            config.write_text(
                json.dumps(
                    {
                        "schema_version": "1.0",
                        "status": "confirmed",
                        "profile_id": "PROFILE-TEST",
                        "confirmed_by": "tester",
                        "confirmed_at": "2026-08-24T12:00:00+08:00",
                        "authority": {
                            "template_path": template.name,
                            "template_sha256": sha256_file(template),
                        },
                        "evidence": {
                            "format_profile_path": evidence.name,
                            "format_profile_sha256": "0" * 64,
                        },
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            with self.assertRaisesRegex(ValueError, "format profile hash"):
                build_docx("# 项目\n", root / "output.docx", "模板测试项目", format_config=config)


if __name__ == "__main__":
    unittest.main()
