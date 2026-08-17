from __future__ import annotations

import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from ingest_scope_items_sqlite import ingest_scope_payload, normalize_payload  # noqa: E402
from extract_xlsx_scope import build_payload, fill_blank_merged_domain_groups  # noqa: E402
from init_project_workbench import initialize_project  # noqa: E402


def sample_payload(rows: list[dict[str, object]]) -> dict[str, object]:
    return {
        "source": {
            "path": "C:/materials/scope.xlsx",
            "size_bytes": 100,
            "sha256": "a" * 64,
        },
        "sheet_count": 1,
        "sheets": [
            {
                "name": "Scope",
                "detected_columns": {
                    "original_name": ["建设内容"],
                    "domain": ["分类"],
                    "item_type": ["费用类型"],
                    "construction_mode": ["建设方式"],
                    "quantity": ["数量"],
                    "unit": ["单位"],
                    "acceptance_target": ["验收目标"],
                },
                "rows": rows,
            }
        ],
    }


class ScopeItemsSqliteTests(unittest.TestCase):
    def test_merged_system_cells_keep_every_module_as_a_scope_item(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            xlsx = Path(tmp) / "scope.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "建设清单"
            sheet.append(["序号", "大类", "系统名称", "模块名称", "备注", "数量"])
            sheet.append([1, "临床业务", "电子病历系统", "临床文书", "利旧升级", 1])
            sheet.append([2, None, None, "病历质控", "利旧升级改成新建", 1])
            sheet.append([None, "说明：本行为横向合并备注，不是建设项", None, None, None, None])
            sheet.merge_cells("B2:B3")
            sheet.merge_cells("C2:C3")
            sheet.merge_cells("B4:F4")
            workbook.save(xlsx)

            payload = build_payload(xlsx)
            items, issues = normalize_payload(payload, "PROJECT-MERGED")

            self.assertEqual({item["standard_name"] for item in items}, {"临床文书", "病历质控"})
            self.assertTrue(all(item["domain"] == "临床业务 / 电子病历系统" for item in items))
            modes = {item["standard_name"]: item["construction_mode"] for item in items}
            self.assertEqual(modes["临床文书"], "upgrade")
            self.assertEqual(modes["病历质控"], "new")
            self.assertEqual(
                len([issue for issue in issues if issue["reason"] == "blank_or_total_row"]),
                1,
            )

    def test_blank_vertical_domain_merge_carries_prior_category_only(self) -> None:
        rows = [
            {2: "大类", 3: "系统名称", 4: "模块名称"},
            {2: "临床业务", 3: "电子病历系统", 4: "临床文书"},
            {2: "", 3: "", 4: "病历质控"},
            {2: "", 3: "", 4: "病案管理"},
        ]
        fill_blank_merged_domain_groups(
            rows,
            ["B3:B4", "C3:C4"],
            ["序号", "大类", "系统名称", "模块名称"],
        )
        self.assertEqual(rows[2][2], "临床业务")
        self.assertEqual(rows[3][2], "临床业务")
        self.assertEqual(rows[2][3], "")
        self.assertEqual(rows[3][3], "")

    def test_normalization_uses_stable_semantic_ids(self) -> None:
        first_rows = [
            {
                "_source_row": 2,
                "建设内容": "1. 电子病历系统",
                "分类": "临床",
                "费用类型": "软件",
                "建设方式": "升级",
                "数量": "2",
                "单位": "套",
                "验收目标": "五级",
            },
            {
                "_source_row": 3,
                "建设内容": "集成平台",
                "分类": "平台",
                "费用类型": "软件",
                "建设方式": "优化",
                "数量": "",
                "单位": "",
                "验收目标": "",
            },
        ]
        second_rows = list(reversed(first_rows))

        first, _ = normalize_payload(sample_payload(first_rows), "PROJECT-1")
        second, _ = normalize_payload(sample_payload(second_rows), "PROJECT-1")

        self.assertEqual(
            {item["scope_id"] for item in first},
            {item["scope_id"] for item in second},
        )
        emr = next(item for item in first if item["standard_name"] == "电子病历系统")
        platform = next(item for item in first if item["standard_name"] == "集成平台")
        self.assertEqual(emr["construction_mode"], "upgrade")
        self.assertEqual(emr["quantity"], 2.0)
        self.assertEqual(platform["construction_mode"], "pending_confirmation")
        self.assertIsNone(platform["quantity"])

        changed_rows = [dict(first_rows[0], **{"建设方式": "利旧升级", "数量": "3"})]
        changed, _ = normalize_payload(sample_payload(changed_rows), "PROJECT-1")
        self.assertEqual(changed[0]["scope_id"], emr["scope_id"])
        self.assertEqual(changed[0]["construction_mode"], "upgrade")

    def test_ingest_is_idempotent_and_preserves_confirmed_fields(self) -> None:
        rows = [
            {
                "_source_row": 2,
                "建设内容": "电子病历系统",
                "分类": "临床",
                "费用类型": "软件",
                "建设方式": "升级",
                "数量": "1",
                "单位": "套",
                "验收目标": "五级",
            }
        ]
        payload = sample_payload(rows)
        with tempfile.TemporaryDirectory() as tmp:
            initialized = initialize_project(
                Path(tmp) / "project", project_code="TEST-SCOPE-001"
            )
            database = Path(initialized["database"])
            first = ingest_scope_payload(
                database, payload, project_code="TEST-SCOPE-001"
            )
            scope_id = first["items"][0]["scope_id"]

            conn = sqlite3.connect(database)
            try:
                conn.execute(
                    """
                    UPDATE project_scope_item
                    SET status='confirmed', standard_name='人工确认名称', construction_mode='reuse'
                    WHERE scope_id=?
                    """,
                    (scope_id,),
                )
                conn.commit()
            finally:
                conn.close()

            second = ingest_scope_payload(
                database, payload, project_code="TEST-SCOPE-001"
            )
            self.assertEqual(first["scope_items_created"], 1)
            self.assertEqual(second["scope_items_created"], 0)
            self.assertEqual(second["scope_items_updated"], 1)

            conn = sqlite3.connect(database)
            try:
                row = conn.execute(
                    "SELECT standard_name,construction_mode,status FROM project_scope_item"
                ).fetchone()
                self.assertEqual(row, ("人工确认名称", "reuse", "confirmed"))
                self.assertEqual(
                    conn.execute("SELECT COUNT(*) FROM evidence_record").fetchone()[0], 1
                )
            finally:
                conn.close()

    def test_merges_exact_semantic_duplicates_but_keeps_row_evidence(self) -> None:
        rows = [
            {
                "_source_row": 2,
                "建设内容": "接口改造",
                "分类": "集成",
                "费用类型": "接口",
                "建设方式": "新建",
                "数量": "",
                "单位": "",
                "验收目标": "",
            },
            {
                "_source_row": 5,
                "建设内容": "接口改造",
                "分类": "集成",
                "费用类型": "接口",
                "建设方式": "新建",
                "数量": "",
                "单位": "",
                "验收目标": "",
            },
        ]
        items, _ = normalize_payload(sample_payload(rows), "PROJECT-1")
        self.assertEqual(len(items), 1)
        self.assertEqual(len(items[0]["source_records"]), 2)


if __name__ == "__main__":
    unittest.main()
