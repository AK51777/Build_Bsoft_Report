from __future__ import annotations

import json
import copy
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from apply_scope_capability_decisions import apply_decisions  # noqa: E402
from build_section_composition_plan import build_composition_plan  # noqa: E402
from build_standard_knowledge_pack import build_pack, chunk_solution  # noqa: E402
from build_standard_knowledge_pack import section_role  # noqa: E402
from export_section_task_packages import export_packages  # noqa: E402
from import_standard_knowledge_pack import import_pack  # noqa: E402
from ingest_scope_items_sqlite import ingest_scope_payload  # noqa: E402
from init_project_workbench import initialize_project  # noqa: E402
from map_scope_capabilities import map_capabilities  # noqa: E402


def make_standard_files(root: Path) -> tuple[Path, Path]:
    solution = root / "company-standard.docx"
    document = Document()
    document.add_heading("建设内容", level=1)
    document.add_heading("电子病历系统", level=2)
    document.add_paragraph(
        "电子病历系统围绕临床文书、病历质控、数据共享和闭环管理形成业务能力。"
        "方案正文只可在客户已确认范围内按项目事实参数化改写，不得反向扩展建设范围。" * 4
    )
    document.add_heading("政策法规", level=1)
    document.add_paragraph(
        "本段仅为历史方案中列示的政策文字，必须重新核验现行状态、官方来源和适用条款。" * 4
    )
    document.save(solution)

    scope = root / "company-scope.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "主表"
    sheet.append(["大类", "系统名称", "模块名称", "产品功能", "三级医院推荐"])
    sheet.append(["临床业务", "电子病历系统", "电子病历系统", "病历质控", "推荐"])
    workbook.save(scope)
    return solution, scope


class StandardKnowledgePackTests(unittest.TestCase):
    def test_short_empty_duplicate_and_short_tail_sections_are_lossless_and_gated(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            solution = root / "lossless-standard.docx"
            document = Document()
            document.add_heading("建设内容", level=1)
            document.add_heading("临床字典管理", level=2)
            document.add_paragraph("短正文必须保留。")
            document.add_heading("仅结构标题", level=2)
            document.add_heading("重复模块", level=2)
            document.add_paragraph("第一次同名路径正文。")
            document.add_heading("重复模块", level=2)
            document.add_paragraph("第二次同名路径正文。")
            document.add_heading("长段拆分", level=2)
            document.add_paragraph("甲" * 80)
            document.add_paragraph("乙" * 80)
            document.add_paragraph("短尾不能丢。")
            document.save(solution)

            scope = root / "scope.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["大类", "系统名称", "模块名称", "产品功能"])
            sheet.append(["平台", "医院信息平台", "临床字典管理", "临床字典管理"])
            workbook.save(scope)

            blocks = chunk_solution(solution, [], min_chars=120, max_chars=100)
            self.assertIn("短正文必须保留。", [block["clean_text"] for block in blocks])
            self.assertTrue(any("短尾不能丢。" in block["clean_text"] for block in blocks))
            structure = next(block for block in blocks if block["heading_path"][-1:] == ["仅结构标题"])
            self.assertEqual(structure["clean_text"], "")
            duplicate_blocks = [block for block in blocks if block["heading_path"][-1:] == ["重复模块"]]
            self.assertEqual(len(duplicate_blocks), 2)
            self.assertNotEqual(duplicate_blocks[0]["source_section_id"], duplicate_blocks[1]["source_section_id"])

            payload = build_pack(solution, scope, title="无损抽取测试")
            coverage = payload["review_summary"]["standard_solution_coverage"]
            self.assertEqual(coverage["status"], "pass")
            self.assertEqual(coverage["source_section_count"], coverage["emitted_section_count"])
            self.assertEqual(coverage["source_heading_count"], coverage["emitted_heading_count"])
            self.assertEqual(
                coverage["short_nonempty_section_count"],
                coverage["short_nonempty_sections_preserved"],
            )
            self.assertEqual(
                coverage["empty_heading_section_count"],
                coverage["empty_heading_sections_preserved"],
            )

            tampered = copy.deepcopy(payload)
            tampered["corpus"]["blocks"] = [
                block
                for block in tampered["corpus"]["blocks"]
                if block["heading_path"][-1:] != ["临床字典管理"]
            ]
            with self.assertRaisesRegex(ValueError, "完整导入门禁"):
                import_pack(root / "tampered.sqlite", tampered)

    def test_build_import_map_and_recall_are_bounded_and_idempotent(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            solution, scope = make_standard_files(root)
            payload = build_pack(solution, scope, title="测试标准知识包")
            serialized = json.dumps(payload, ensure_ascii=False)
            self.assertNotIn(str(root), serialized)
            self.assertEqual(payload["permission_scope"], "internal_company_reuse")
            self.assertEqual(section_role(["政策法规"]), "policy")
            policy_blocks = [
                block for block in payload["corpus"]["blocks"]
                if block["section_role"] == "policy"
            ]
            self.assertTrue(policy_blocks)
            self.assertTrue(all(block["review_status"] == "prohibited" for block in policy_blocks))
            self.assertTrue(payload["capabilities"][0]["standard_block_ids"])
            capability = payload["capabilities"][0]
            capability["standard_block_relations"] = [
                {
                    "block_id": capability["standard_block_ids"][0],
                    "relation_type": "standard_description",
                    "priority": 1,
                    "review_status": "approved",
                    "root_heading_path": ["建设内容", "电子病历系统"],
                    "relation_order": 1,
                    "verbatim_eligible": True,
                }
            ]

            initialized = initialize_project(
                root / "project", project_code="PACK-TEST-001",
                official_name="标准知识包测试项目", owner_name="测试医院",
            )
            database = Path(initialized["database"])
            first = import_pack(database, payload)
            second = import_pack(database, payload)
            self.assertEqual(first["blocks_imported"], second["blocks_imported"])
            self.assertEqual(first["capability_block_relations_imported"], 1)
            conn = sqlite3.connect(database)
            try:
                relation = conn.execute(
                    """
                    SELECT root_heading_path_json,relation_order,verbatim_eligible
                    FROM capability_solution_block_relation
                    """
                ).fetchone()
            finally:
                conn.close()
            self.assertEqual(json.loads(relation[0]), ["建设内容", "电子病历系统"])
            self.assertEqual(relation[1:], (1, 1))

            ingest_scope_payload(
                database,
                {
                    "source": {"path": "scope.xlsx", "sha256": "a" * 64},
                    "sheets": [{
                        "name": "建设清单",
                        "detected_columns": {"original_name": ["建设内容"]},
                        "rows": [{"_source_row": 2, "建设内容": "电子病历系统"}],
                    }],
                },
                project_code="PACK-TEST-001",
            )
            conn = sqlite3.connect(database)
            try:
                conn.execute("UPDATE project_scope_item SET status='confirmed'")
                conn.commit()
            finally:
                conn.close()

            mapping = map_capabilities(database, "PACK-TEST-001")
            self.assertEqual(len(mapping["candidates"]), 1)
            apply_decisions(
                database,
                {
                    "project_code": "PACK-TEST-001",
                    "reviewed_by": "reviewer",
                    "decisions": [{
                        "map_id": mapping["candidates"][0]["map_id"],
                        "decision": "confirmed",
                        "review_note": "客户范围与标准能力一致",
                    }],
                },
            )
            plan_result = build_composition_plan(database, "PACK-TEST-001")
            self.assertEqual(plan_result["not_applicable_count"], 3)
            construction = next(
                item for item in plan_result["plans"] if item["chapter_code"] == "5.1.1"
            )
            self.assertNotIn("capability_mapping_review", construction["missing_source_types"])
            self.assertIn("capability_block_scope_review", construction["missing_source_types"])
            self.assertGreaterEqual(construction["outline_node_count"], 3)
            self.assertEqual(construction["outline_heading_levels"]["4"], 1)
            self.assertEqual(construction["outline_heading_levels"]["5"], 1)

            conn = sqlite3.connect(database)
            try:
                plan_id = conn.execute(
                    "SELECT plan_id FROM section_composition_plan WHERE project_id=(SELECT project_id FROM project WHERE project_code=?) AND chapter_code='5.1.1'",
                    ("PACK-TEST-001",),
                ).fetchone()[0]
                source_types = conn.execute(
                    "SELECT source_type,usage_mode FROM section_plan_source WHERE plan_id=?",
                    (plan_id,),
                ).fetchall()
                self.assertIn(("capability", "parameterized"), source_types)
                self.assertIn(("corpus", "structure_only"), source_types)
                planned_block_ids = {
                    row[0]
                    for row in conn.execute(
                        "SELECT source_object_id FROM section_plan_source "
                        "WHERE plan_id=? AND source_type='corpus'",
                        (plan_id,),
                    )
                }
                outline_block_ids = {
                    row[0]
                    for row in conn.execute(
                        "SELECT source_object_id FROM section_outline_node "
                        "WHERE plan_id=? AND source_type='corpus'",
                        (plan_id,),
                    )
                }
                self.assertEqual(planned_block_ids, outline_block_ids)
                self.assertEqual(conn.execute("SELECT COUNT(*) FROM project_scope_item").fetchone()[0], 1)
            finally:
                conn.close()

            packages = root / "packages"
            exported = export_packages(database, "PACK-TEST-001", packages)
            self.assertEqual(exported["package_count"], 25)
            package = json.loads(
                next(packages.glob("CH5.1.1-*.json")).read_text(encoding="utf-8")
            )
            self.assertGreaterEqual(len(package["outline_nodes"]), 3)
            self.assertIn("heading_path_json", next(
                source["data"] for source in package["sources"]
                if source["source_type"] == "corpus"
            ))


if __name__ == "__main__":
    unittest.main()
