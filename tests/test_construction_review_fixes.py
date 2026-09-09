from __future__ import annotations

import csv
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import test_construction_workflow as fixtures
from build_report_docx import parse_table_row
from construction_word import audit_word
from construction_workflow import source_payload, read, confirm_and_generate


class ScopeTableFidelityTests(unittest.TestCase):
    def setUp(self):
        self.fixture = fixtures.ConstructionWorkflowTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)

    def test_table_delimiters_preserve_empty_cells_and_literal_pipes(self):
        cases = [
            ("|2|管理数据中心（ODR)|||", ["2", "管理数据中心（ODR)", "", ""]),
            ("||业务||模块||", ["", "业务", "", "模块", ""]),
            ("|||||", ["", "", "", ""]),
            (r"|1|甲\|乙||", ["1", "甲|乙", ""]),
            (r"|1|末尾\|", ["1", "末尾|"]),
            (r"|1|路径\\|末列|", ["1", "路径\\\\", "末列"]),
        ]
        for text, expected in cases:
            with self.subTest(text=text):
                self.assertEqual(parse_table_row(text), expected)

    def test_pasted_markdown_with_empty_edge_columns_imports(self):
        source = self.fixture.root / "scope.md"
        source.write_text(
            "|序号|大类|系统名称|模块名称|\n|---|---|---|---|\n"
            "|1|临床业务|电子病历系统||\n||临床业务|未知系统||\n",
            encoding="utf-8",
        )
        sheet = source_payload(source)["sheets"][0]
        self.assertEqual([[row[h] for h in sheet["headers"]] for row in sheet["rows"]],
                         [["1", "临床业务", "电子病历系统", ""], ["", "临床业务", "未知系统", ""]])

    def build_edge_table(self):
        headers = ["序号", "大类", "系统名称", "模块名称", "备注", "数量", "预留"]
        rows = [["", "临床业务", "电子病历系统", "电子病历系统", "甲|乙\n第二行", "0", ""],
                ["2", "临床业务", "未知系统", "未知模块", "", "", ""]]
        with self.fixture.scope.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(headers)
            writer.writerows(rows)
        result = self.fixture.confirm(self.fixture.prepare())
        manifest = read(self.fixture.work / "运行数据/construction-assembly-manifest.json")
        return Path(result["docx_path"]), manifest, [headers, *rows]

    def test_csv_to_word_preserves_snapshot_matrix_including_empty_cells(self):
        output, manifest, expected = self.build_edge_table()
        actual = [[cell.text for cell in row.cells] for row in Document(output).tables[0].rows]
        self.assertEqual(actual, expected)
        self.assertTrue(audit_word(output, manifest)["valid"])

    def test_shifted_word_row_fails_even_if_shared_parser_has_same_bug(self):
        output, manifest, expected = self.build_edge_table()
        document = Document(output)
        shifted = expected[1][1:] + [""]
        for cell, text in zip(document.tables[0].rows[1].cells, shifted):
            cell.text = text
        document.save(output)

        def old_parser(line):
            # The old shared transformation shifts leading blanks and loses tails.
            return [cell.strip() for cell in line.strip().strip("|").split("|")]

        with patch("build_report_docx.parse_table_row", side_effect=old_parser):
            result = audit_word(output, manifest)
        self.assertFalse(result["valid"])
        self.assertIsNotNone(result["first_mismatch"])


class ParentOverrideTests(unittest.TestCase):
    def setUp(self):
        self.fixture = fixtures.ConstructionWorkflowTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.status = self.fixture.prepare()
        self.state_path = self.fixture.work / "运行数据/construction-state.json"
        self.review = read(read(self.state_path)["review_path"])
        self.first = self.review["items"][0]

    def confirm_override(self, override):
        return confirm_and_generate(self.fixture.work, {
            "review_id": self.status["review_id"], "reviewed_by": "test-user",
            "user_reply": "其余按建议，指定项目按本次明确修改处理", "accept_all": True,
            "overrides": [override],
        })

    def current_artifacts(self):
        state = read(self.state_path)
        return state, read(state["confirmation_path"]), read(state["manifest_path"])

    def test_parent_only_override_keeps_candidate_and_can_retry(self):
        override = {"scope_row_id": self.first["scope_row_id"], "parent_path": ["新业务"]}
        result = self.confirm_override(override)
        state, receipt, manifest = self.current_artifacts()
        decision = receipt["decisions"][0]
        self.assertEqual(decision["candidate_id"], self.first["proposal"]["candidate_id"])
        self.assertEqual(decision["decision"], "confirmed")
        self.assertEqual(manifest["application_software_solution"]["items"][0]["display_parent_path"], ["新业务"])
        self.assertEqual(receipt["decisions"][1]["decision"], "confirmed_gap")
        self.assertTrue(audit_word(result["docx_path"], manifest)["valid"])
        before = Path(result["docx_path"]).stat().st_mtime_ns
        with patch("construction_workflow.generate_word", side_effect=AssertionError("unexpected rebuild")):
            self.confirm_override(override)
        self.assertEqual(Path(result["docx_path"]).stat().st_mtime_ns, before)

    def test_gap_override_clears_inherited_selection(self):
        self.confirm_override({"scope_row_id": self.first["scope_row_id"], "decision": "confirmed_gap"})
        _, receipt, manifest = self.current_artifacts()
        decision = receipt["decisions"][0]
        self.assertEqual(decision["decision"], "confirmed_gap")
        for field in ("candidate_id", "capability_id", "root_heading_path"):
            self.assertNotIn(field, decision)
        item = manifest["application_software_solution"]["items"][0]
        self.assertEqual(item["status"], "pending_supplement")
        self.assertEqual(item["fragments"], [])

    def test_manual_capability_does_not_inherit_candidate_selector(self):
        candidate = next(c for c in self.first["candidates"]
                         if c["candidate_id"] == self.first["proposal"]["candidate_id"])
        self.confirm_override({"scope_row_id": self.first["scope_row_id"],
                               "capability_id": candidate["capability_id"],
                               "root_heading_path": candidate["root_heading_path"]})
        _, receipt, manifest = self.current_artifacts()
        decision = receipt["decisions"][0]
        self.assertNotIn("candidate_id", decision)
        self.assertEqual(decision["capability_id"], candidate["capability_id"])
        self.assertEqual(manifest["application_software_solution"]["items"][0]["capability_id"], candidate["capability_id"])

    def test_candidate_only_override_keeps_confirmed_decision(self):
        self.confirm_override({"scope_row_id": self.first["scope_row_id"],
                               "candidate_id": self.first["proposal"]["candidate_id"]})
        _, receipt, _ = self.current_artifacts()
        self.assertEqual(receipt["decisions"][0]["decision"], "confirmed")

    def test_candidate_from_another_row_is_still_rejected(self):
        with self.assertRaisesRegex(ValueError, "candidate does not belong to row/run"):
            self.confirm_override({"scope_row_id": self.review["items"][1]["scope_row_id"],
                                   "candidate_id": self.first["proposal"]["candidate_id"]})


if __name__ == "__main__":
    unittest.main()
