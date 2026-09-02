from __future__ import annotations

import sys
import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from build_policy_catalog import build_catalog  # noqa: E402
from import_verified_policies_postgres import validate_policies  # noqa: E402
from import_policy_catalog_postgres import normalize_catalog, validate_catalog  # noqa: E402
from import_standard_knowledge_pack_postgres import (  # noqa: E402
    delete_existing_document_blocks,
    retire_superseded_packages,
)
from postgres_knowledge_db import (  # noqa: E402
    migration_dir,
    migration_hash_is_accepted,
    validate_schema,
)
from sync_postgres_knowledge_snapshot import (  # noqa: E402
    build_pack_snapshot,
    catalog_payload_from_rows,
    record_snapshot,
)
from init_project_workbench import initialize_project  # noqa: E402


def make_policy_catalog(path: Path) -> None:
    workbook = Workbook()
    cover = workbook.active
    cover.title = "首页Coverpage"
    cover["A1"] = "医疗卫生政策索引库"
    current = workbook.create_sheet("国家与部委政策CN")
    current.append([])
    current.append(
        [None, "级别", "类别", "关键词", "字号", "标题", "发文日期", "发布部门与机构", "文件数", "索引号", "备注", "外部链接"]
    )
    current.append([None, " (CN9) 九、互联网与新技术应用 INN", None, None, None, None, None, None, None, None, None, None])
    current.append([None, "国家", "新技术应用", "AI人工智能/数据运用", "国办发〔2025〕1号", "《医疗人工智能应用指导意见》", "2025-05-01", "国务院办公厅", 1, "CN9.001", "", "https://example.gov.cn/policy/1"])
    current.append([None, "部委", "信息化与统计上报", "信息化", "", "《医疗卫生信息化建设指南》", "2025-06-01", "国家卫生健康委", 1, "CN9.002", "待核验有效性", "https://example.gov.cn/policy/2"])
    old = workbook.create_sheet("国家与部委政策CN V1(隐藏)")
    old.sheet_state = "hidden"
    old.append([])
    old.append([None, "级别", "类别", "关键词", "字号", "标题", "发文日期", "发布部门与机构", "文件数", "索引号", "备注", "外部链接"])
    old.append([None, "国家", "旧版", "旧版", "", "《旧版政策》", "2020-01-01", "国务院", 1, "OLD.001", "", ""])
    workbook.save(path)


class PostgresKnowledgeRepositoryTests(unittest.TestCase):
    def test_legacy_standard_solution_snapshot_without_coverage_proof_is_blocked(self) -> None:
        class Cursor:
            def __init__(self, connection):
                self.connection = connection
                self.description = []
                self.rows = []

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def execute(self, sql, _params=None):
                for token, columns, rows in self.connection.responses:
                    if token in sql:
                        self.description = [
                            type("Column", (), {"name": column})() for column in columns
                        ]
                        self.rows = rows
                        return
                raise AssertionError(sql)

            def fetchall(self):
                return self.rows

        class Connection:
            def __init__(self, responses):
                self.responses = responses

            def cursor(self):
                return Cursor(self)

        block_columns = [
            "block_id", "source_location", "heading_path", "section_role", "module_code",
            "clean_text", "reuse_class", "quality_level", "prerequisites", "variable_slots",
            "forbidden_terms", "length_band", "text_hash", "content_type", "semantic_section",
            "content_slot", "source_order", "block_index", "adaptation_mode", "assessment_targets",
            "construction_scope_tags", "content_format", "content_payload", "asset_manifest",
            "visible_text_hash",
        ]
        connection = Connection([
            ("runtime_package_source", ["source_role", "file_name", "source_sha256"], [("standard_solution", "standard.docx", "a" * 64)]),
            ("runtime_corpus_document", ["corpus_document_id", "document_type", "project_type", "quality_level", "permission_scope", "version", "source_corpus_type"], [("DOC-1", "feasibility_study", "hospital_informationization", "B", "internal_company_reuse", "1", "legacy_unspecified")]),
            ("runtime_corpus_block", block_columns, [("BLOCK-1", "paragraph:7", ["系统", "模块"], "construction", "", "正文", "B", "B", [], [], [], {}, "b" * 64, "legacy_unspecified", "", "", 0, 7, "structure_only", [], [], "plain_text", {}, [], "b" * 64)]),
            ("runtime_product_capability", ["capability_id", "product_code", "product_name", "capability_name", "capability_description", "category", "module_name", "selection_rules", "prerequisites", "interface_dependencies", "exclusions", "applicable_versions", "source_location"], [("CAP-1", "P1", "系统", "模块", "", "", "模块", [], [], [], [], [], "row:1")]),
            (
                "runtime_capability_block",
                [
                    "capability_id", "block_id", "relation_type", "priority",
                    "review_status", "root_heading_path", "relation_order",
                    "verbatim_eligible",
                ],
                [
                    (
                        "CAP-1", "BLOCK-1", "standard_description", 1,
                        "approved", ["系统", "模块"], 1, True,
                    )
                ],
            ),
        ])
        with self.assertRaisesRegex(ValueError, "完整导入门禁"):
            build_pack_snapshot(
                connection,
                "medical_report_kb",
                {
                    "package_id": "PACK-1",
                    "schema_version": "1.0",
                    "title": "标准包",
                    "permission_scope": "internal_company_reuse",
                    "review_summary": {},
                    "content_hash": "c" * 64,
                },
            )

    def test_policy_catalog_build_is_stable_and_uses_current_visible_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "policy-index.xlsx"
            make_policy_catalog(workbook_path)
            first = build_catalog(workbook_path)
            second = build_catalog(workbook_path)
            self.assertEqual(first["catalog_id"], second["catalog_id"])
            self.assertEqual(first["content_hash"], second["content_hash"])
            self.assertEqual(first["worksheet_name"], "国家与部委政策CN")
            self.assertEqual(len(first["records"]), 2)
            self.assertEqual(first["records"][0]["catalog_group_code"], "CN9")
            self.assertEqual(first["records"][0]["keyword_tags"], ["AI人工智能", "数据运用"])
            self.assertEqual(first["records"][0]["publish_date"], "2025-05-01")
            self.assertNotIn(str(Path(tmp)), str(first))
            validate_catalog(first)

    def test_policy_catalog_normalization_preserves_duplicate_index_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "policy-index.xlsx"
            make_policy_catalog(workbook_path)
            payload = build_catalog(workbook_path)
            payload["records"][1]["source_index_no"] = payload["records"][0]["source_index_no"]
            payload["records"][1]["catalog_entry_id"] = payload["records"][0]["catalog_entry_id"]
            normalized = normalize_catalog(payload)
            self.assertEqual(
                [record["index_occurrence"] for record in normalized["records"]], [1, 2]
            )
            self.assertTrue(all(record["index_conflict"] for record in normalized["records"]))
            self.assertEqual(
                len({record["catalog_entry_id"] for record in normalized["records"]}), 2
            )
            self.assertTrue(
                all(record["jurisdiction_level"] == "national" for record in normalized["records"])
            )
            self.assertTrue(
                all(record["jurisdiction_code"] == "100000" for record in normalized["records"])
            )
            validate_catalog(payload)

    def test_policy_catalog_validation_rejects_duplicate_source_row(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "policy-index.xlsx"
            make_policy_catalog(workbook_path)
            payload = build_catalog(workbook_path)
            payload["records"][1]["source_row"] = payload["records"][0]["source_row"]
            with self.assertRaisesRegex(ValueError, "source_row"):
                validate_catalog(payload)

    def test_postgres_schema_contains_runtime_boundaries(self) -> None:
        schema_sql = (migration_dir() / "002_shared_knowledge_policy_schema.sql").read_text(encoding="utf-8")
        for table in (
            "knowledge_package",
            "corpus_block",
            "product_capability",
            "capability_block",
            "policy_catalog_entry",
            "policy_document",
            "policy_clause",
        ):
            self.assertIn(f"medical_report_kb.{table}", schema_sql)
        self.assertIn("runtime_corpus_block", schema_sql)
        self.assertIn("runtime_policy_clause", schema_sql)
        self.assertIn("policy.verification_status = 'verified'", schema_sql)
        self.assertIn("clause.review_status = 'approved'", schema_sql)

    def test_policy_catalog_runtime_view_is_refreshed_after_conflict_columns(self) -> None:
        migration_sql = (
            migration_dir() / "005_policy_catalog_runtime_view.sql"
        ).read_text(encoding="utf-8")
        self.assertIn("CREATE OR REPLACE VIEW medical_report_kb.runtime_policy_catalog_entry", migration_sql)
        self.assertIn("SELECT entry.*", migration_sql)
        catalog_view_sql = (
            migration_dir() / "006_runtime_policy_catalog.sql"
        ).read_text(encoding="utf-8")
        self.assertIn("CREATE OR REPLACE VIEW medical_report_kb.runtime_policy_catalog", catalog_view_sql)

    def test_policy_catalog_jurisdiction_migration_is_runtime_visible(self) -> None:
        migration_sql = (
            migration_dir() / "011_policy_catalog_jurisdiction.sql"
        ).read_text(encoding="utf-8")
        for column in ("jurisdiction_level", "jurisdiction_code", "jurisdiction_name"):
            self.assertIn(column, migration_sql)
        self.assertIn("authority_level_label IN", migration_sql)
        self.assertIn("runtime_policy_catalog_entry", migration_sql)

    def test_capability_match_scope_migration_appends_view_column(self) -> None:
        migration_sql = (
            migration_dir() / "007_capability_block_match_scope.sql"
        ).read_text(encoding="utf-8")
        self.assertNotIn("capability.*", migration_sql)
        self.assertIn(
            "package.content_hash AS package_content_hash,\n  capability.block_match_scope",
            migration_sql,
        )

    def test_construction_subtree_migration_preserves_capability_view_columns(self) -> None:
        migration_sql = (
            migration_dir() / "009_construction_solution_subtree.sql"
        ).read_text(encoding="utf-8")
        self.assertIn(
            "relation.priority,\n  relation.review_status,\n  relation.root_heading_path",
            migration_sql,
        )

    def test_standard_solution_coverage_migration_appends_runtime_view_columns(self) -> None:
        migration_sql = (
            migration_dir() / "010_standard_solution_coverage.sql"
        ).read_text(encoding="utf-8")
        self.assertIn(
            "block.visible_text_hash,\n  block.source_section_id,\n"
            "  block.chunk_index,\n  block.source_is_heading",
            migration_sql,
        )
        self.assertNotIn(
            "block.block_index,\n  block.source_section_id",
            migration_sql,
        )

    def test_policy_catalog_snapshot_payload_preserves_candidate_boundaries(self) -> None:
        payload = catalog_payload_from_rows(
            {
                "schema_version": "1.0",
                "catalog_id": "CAT-001",
                "catalog_scope": "department_policy_index",
                "title": "部门政策目录",
                "permission_scope": "internal_company_reference",
                "source_file_name": "policy-index.xlsx",
                "source_sha256": "a" * 64,
                "worksheet_name": "当前政策",
                "content_hash": "b" * 64,
                "metadata": {"built_at": "2026-08-14T00:00:00+00:00"},
            },
            [
                {
                    "catalog_entry_id": "ENTRY-001",
                    "source_row": 10,
                    "source_index_no": "CN6.061",
                    "index_occurrence": 1,
                    "index_conflict": True,
                    "identity_key": "identity",
                    "catalog_group_code": "CN6",
                    "catalog_group_name": "卫生健康",
                    "authority_level_label": "国家",
                    "category_name": "政策",
                    "keyword_text": "医疗信息化",
                    "keyword_tags": ["医疗信息化"],
                    "document_no": "",
                    "title": "测试政策",
                    "publish_date": None,
                    "publish_date_raw": "",
                    "issuer": "测试单位",
                    "file_count": 1,
                    "notes": "",
                    "external_url": "",
                    "verification_status": "unverified",
                    "entry_status": "active",
                    "row_hash": "c" * 64,
                }
            ],
        )
        validate_catalog(payload)
        self.assertEqual(payload["records"][0]["source_row"], 10)
        self.assertTrue(payload["records"][0]["index_conflict"])
        self.assertEqual(payload["records"][0]["verification_status"], "unverified")

    def test_schema_name_is_restricted(self) -> None:
        self.assertEqual(validate_schema("medical_report_kb"), "medical_report_kb")
        with self.assertRaises(ValueError):
            validate_schema("medical_report_kb; DROP SCHEMA public")

    def test_only_exact_legacy_migration_hash_pair_is_accepted(self) -> None:
        version = "002_shared_knowledge_policy_schema"
        canonical = "2b0015882d48b9ce2977e7ed079b8eae7d97b31a29d8ef889a8adbc00913b86a"
        legacy = "1e64be3aff8afc3156d303f6d86e6757fa44fc50ae00421447e948fcaa549e1e"
        self.assertTrue(migration_hash_is_accepted(version, canonical, canonical))
        self.assertTrue(migration_hash_is_accepted(version, legacy, canonical))
        self.assertFalse(migration_hash_is_accepted(version, "f" * 64, canonical))
        self.assertFalse(migration_hash_is_accepted(version, legacy, "e" * 64))
        self.assertFalse(migration_hash_is_accepted("003_runtime_snapshot_views", legacy, canonical))

    def test_publishing_standard_pack_retires_same_source_family(self) -> None:
        class Cursor:
            rowcount = 2

            def execute(self, statement, params):
                self.statement = statement
                self.params = params

        cursor = Cursor()
        count = retire_superseded_packages(
            cursor,
            schema="medical_report_kb",
            package_id="PACK-NEW",
            solution_source_id="SOURCE-SOLUTION",
            publish=True,
        )
        self.assertEqual(count, 2)
        self.assertIn("package_status='retired'", cursor.statement)
        self.assertIn("source_role='standard_solution'", cursor.statement)
        self.assertEqual(cursor.params, ("PACK-NEW", "SOURCE-SOLUTION"))

    def test_rebuilt_document_replaces_old_blocks_before_insert(self) -> None:
        class Cursor:
            rowcount = 2364

            def execute(self, statement, params):
                self.statement = statement
                self.params = params

        cursor = Cursor()
        count = delete_existing_document_blocks(
            cursor,
            schema="medical_report_kb",
            corpus_document_id="STDDOC-STABLE",
        )
        self.assertEqual(count, 2364)
        self.assertIn("DELETE FROM medical_report_kb.corpus_block", cursor.statement)
        self.assertEqual(cursor.params, ("STDDOC-STABLE",))

    def test_reviewed_standard_pack_does_not_retire_published_family(self) -> None:
        class Cursor:
            def execute(self, *_args):
                raise AssertionError("reviewed imports must not retire published packages")

        self.assertEqual(
            retire_superseded_packages(
                Cursor(),
                schema="medical_report_kb",
                package_id="PACK-REVIEW",
                solution_source_id="SOURCE-SOLUTION",
                publish=False,
            ),
            0,
        )

    def test_published_policy_requires_verified_document_and_clause(self) -> None:
        payload = {
            "policies": [
                {
                    "title": "测试政策",
                    "issuer": "测试机关",
                    "official_url": "https://example.gov.cn/policy",
                    "verification_status": "verified",
                    "clauses": [
                        {
                            "original_text": "支持医疗机构依法依规推进信息化建设。",
                            "verification_status": "unverified",
                        }
                    ],
                }
            ]
        }
        with self.assertRaisesRegex(ValueError, "clause is not verified"):
            validate_policies(payload, publish=True)

    def test_published_policy_requires_safe_summary_and_section_permissions(self) -> None:
        policy = {
            "title": "测试政策",
            "issuer": "测试机关",
            "official_url": "https://example.gov.cn/policy",
            "verification_status": "verified",
            "clauses": [
                {
                    "original_text": "支持医疗机构依法依规推进信息化建设。",
                    "normalized_summary": "",
                    "permitted_sections": ["basis", "policy_background"],
                    "verification_status": "verified",
                }
            ],
        }
        with self.assertRaisesRegex(ValueError, "normalized_summary is missing"):
            validate_policies({"policies": [policy]}, publish=True)
        policy["clauses"][0]["normalized_summary"] = "支持依法依规推进信息化建设。"
        policy["clauses"][0]["permitted_sections"] = []
        with self.assertRaisesRegex(ValueError, "permitted_sections is missing"):
            validate_policies({"policies": [policy]}, publish=True)


    def test_sqlite_snapshot_is_idempotent_and_hash_bound(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            initialized = initialize_project(
                root / "project",
                project_code="SNAPSHOT-001",
                official_name="知识快照测试项目",
            )
            database = Path(initialized["database"])
            first = record_snapshot(
                database,
                "SNAPSHOT-001",
                source_type="knowledge_package",
                source_id="PACK-001",
                content_hash="a" * 64,
                server_schema="medical_report_kb",
                items=[("corpus_block", "BLOCK-001", "b" * 64, {"clean_text": "测试内容"})],
                metadata={"blocks": 1},
            )
            second = record_snapshot(
                database,
                "SNAPSHOT-001",
                source_type="knowledge_package",
                source_id="PACK-001",
                content_hash="a" * 64,
                server_schema="medical_report_kb",
                items=[("corpus_block", "BLOCK-001", "b" * 64, {"clean_text": "测试内容"})],
                metadata={"blocks": 1},
            )
            catalog_snapshot = record_snapshot(
                database,
                "SNAPSHOT-001",
                source_type="policy_catalog",
                source_id="CAT-001",
                content_hash="c" * 64,
                server_schema="medical_report_kb",
                items=[
                    (
                        "policy_catalog_entry",
                        "ENTRY-001",
                        "d" * 64,
                        {"title": "候选政策", "candidate_only": True},
                    )
                ],
                metadata={"records": 1, "candidate_only": True},
            )
            self.assertEqual(first, second)
            self.assertNotEqual(first, catalog_snapshot)
            connection = sqlite3.connect(database)
            try:
                self.assertEqual(
                    connection.execute("SELECT COUNT(*) FROM shared_knowledge_snapshot").fetchone()[0],
                    2,
                )
                self.assertEqual(
                    connection.execute("SELECT COUNT(*) FROM shared_knowledge_snapshot_item").fetchone()[0],
                    2,
                )
                self.assertEqual(
                    connection.execute(
                        "SELECT COUNT(*) FROM shared_knowledge_snapshot_item WHERE item_type='policy_catalog_entry'"
                    ).fetchone()[0],
                    1,
                )
            finally:
                connection.close()


if __name__ == "__main__":
    unittest.main()
