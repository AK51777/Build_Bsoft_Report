from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.style import WD_STYLE_TYPE


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from knowledge_db import sha256_file  # noqa: E402
from medical_report_mcp_server import (  # noqa: E402
    MedicalReportMCP,
    ServerConfig,
    TOOLS,
)


class MedicalReportMCPTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tempdir = tempfile.TemporaryDirectory()
        self.root = Path(self.tempdir.name).resolve()
        self.database = self.root / "knowledge.sqlite"
        self.database.write_bytes(b"sqlite-placeholder")
        self.xlsx = self.root / "scope.xlsx"
        self.xlsx.write_bytes(b"xlsx-placeholder")
        self.server = MedicalReportMCP(
            ServerConfig((self.root,), response_mode="review_metadata")
        )

    def tearDown(self) -> None:
        self.tempdir.cleanup()

    def test_initialize_and_tool_discovery(self) -> None:
        initialized = self.server.dispatch(
            {
                "jsonrpc": "2.0",
                "id": 1,
                "method": "initialize",
                "params": {"protocolVersion": "2025-06-18"},
            }
        )
        self.assertEqual(initialized["result"]["protocolVersion"], "2025-06-18")
        self.assertIn("tools", initialized["result"]["capabilities"])

        listed = self.server.dispatch(
            {"jsonrpc": "2.0", "id": 2, "method": "tools/list", "params": {}}
        )
        names = {tool["name"] for tool in listed["result"]["tools"]}
        self.assertEqual(names, {tool["name"] for tool in TOOLS})
        self.assertEqual(
            names,
            {
                "service_status",
                "construction_prepare_review",
                "construction_apply_and_assemble",
                "word_generate",
            },
        )

    def test_stdio_protocol_is_utf8_on_windows_and_other_platforms(self) -> None:
        config = self.root / "mcp-config.json"
        config.write_text(
            json.dumps(
                {
                    "schema_version": "1.0",
                    "allowed_project_roots": [str(self.root)],
                    "response_mode": "review_metadata",
                    "allow_delivery_mode": False,
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        requests = (
            '{"jsonrpc":"2.0","id":1,"method":"tools/list","params":{}}\n'
        ).encode("utf-8")
        completed = subprocess.run(
            [
                sys.executable,
                str(SCRIPTS / "medical_report_mcp_server.py"),
                "--config",
                str(config),
            ],
            input=requests,
            capture_output=True,
            check=False,
        )

        self.assertEqual(completed.returncode, 0, completed.stderr.decode("utf-8"))
        decoded = completed.stdout.decode("utf-8")
        response = json.loads(decoded)
        description = next(
            tool["description"]
            for tool in response["result"]["tools"]
            if tool["name"] == "construction_prepare_review"
        )
        self.assertIn("建设清单", description)

    def test_paths_cannot_escape_project_root(self) -> None:
        with tempfile.TemporaryDirectory() as other:
            outside = Path(other) / "outside.xlsx"
            outside.write_bytes(b"outside")
            response = self.server.dispatch(
                {
                    "jsonrpc": "2.0",
                    "id": 3,
                    "method": "tools/call",
                    "params": {
                        "name": "construction_prepare_review",
                        "arguments": {
                            "project_root": str(self.root),
                            "database_path": str(self.database),
                            "scope_xlsx_path": str(outside),
                            "project_code": "TEST-001",
                        },
                    },
                }
            )
        result = response["result"]
        self.assertTrue(result["isError"])
        self.assertEqual(result["structuredContent"]["error_type"], "PermissionError")

    @patch("medical_report_mcp_server.match_scope")
    @patch("medical_report_mcp_server.capture_scope_snapshot")
    @patch("medical_report_mcp_server.ingest_scope_payload")
    @patch("medical_report_mcp_server.build_scope_payload")
    def test_prepare_review_returns_only_candidate_metadata(
        self,
        build_scope,
        ingest_scope,
        capture_scope,
        match_scope,
    ) -> None:
        scope_payload = {
            "source": {"path": str(self.xlsx), "sha256": "scope-hash"},
            "sheet_count": 1,
            "sheets": [],
        }
        build_scope.return_value = scope_payload
        ingest_scope.return_value = {"scope_items_total_in_input": 1, "items": []}
        capture_scope.return_value = {
            "scope_snapshot_id": "SNAPSHOT-001",
            "scope_row_count": 1,
        }
        match_scope.return_value = {
            "match_run_id": "MATCH-001",
            "summary": {
                "scope_rows": 1,
                "auto_confirmed_exact": 0,
                "needs_human_review": 1,
                "content_missing": 0,
            },
            "items": [
                {
                    "scope_row_id": "ROW-001",
                    "source_ordinal": 1,
                    "original_name": "患者全息视图（移动端）",
                    "hierarchy": ["临床应用", "患者全息视图（移动端）"],
                    "state": "needs_human_review",
                    "question": "是否对应标准模块？",
                    "candidates": [
                        {
                            "candidate_id": "CANDIDATE-001",
                            "product_name": "基于CDR的医院临床应用",
                            "module_name": "移动患者全息视图",
                            "match_class": "similar",
                            "name_score": 0.9,
                            "parent_score": 1.0,
                            "root_heading_path": ["建设内容", "移动患者全息视图"],
                            "candidate_status": "ready",
                            "reason": "名称相似",
                            "subtree_block_ids": ["BLOCK-SECRET"],
                        }
                    ],
                }
            ],
        }

        result = self.server.construction_prepare_review(
            {
                "project_root": str(self.root),
                "database_path": str(self.database),
                "scope_xlsx_path": str(self.xlsx),
                "project_code": "TEST-001",
                "package_id": "PACK-001",
            }
        )

        self.assertEqual(result["status"], "needs_human_confirmation")
        self.assertEqual(result["match_run_id"], "MATCH-001")
        self.assertEqual(len(result["review_items"]), 1)
        self.assertNotIn("subtree_block_ids", result["review_items"][0]["candidates"][0])
        self.assertFalse(any("BLOCK-SECRET" in json.dumps(value) for value in result.values()))
        self.assertTrue(Path(result["artifacts"]["match_review_json"]["path"]).is_file())

    @patch("medical_report_mcp_server.validate_manifest")
    @patch("medical_report_mcp_server.assemble")
    @patch("medical_report_mcp_server.apply_decisions")
    def test_apply_and_assemble_reports_preview_gate(
        self,
        apply_decisions,
        assemble,
        validate_manifest,
    ) -> None:
        apply_decisions.return_value = {"match_run_id": "MATCH-001", "applied": 1}
        manifest = {
            "manifest_id": "MANIFEST-001",
            "status": "blocked",
            "preview_only": True,
            "construction_list_import": {"display_payload": {"sheets": []}},
            "application_software_solution": {
                "items": [
                    {
                        "source_ordinal": 1,
                        "original_name": "新增模块",
                        "status": "pending_supplement",
                        "fragments": [],
                    }
                ]
            },
        }
        assemble.return_value = (manifest, "## 建设清单\n\n## 应用软件建设方案\n")
        validate_manifest.return_value = {
            "valid": False,
            "issues": [{"code": "unresolved_working_preview", "blocking": True}],
        }

        result = self.server.construction_apply_and_assemble(
            {
                "project_root": str(self.root),
                "database_path": str(self.database),
                "project_code": "TEST-001",
                "match_run_id": "MATCH-001",
                "reviewed_by": "human-reviewer",
                "decisions": [{"source_ordinal": 1, "decision": "confirmed_gap"}],
                "allow_unresolved_preview": True,
            }
        )

        self.assertEqual(result["status"], "blocked_working_preview")
        self.assertFalse(result["validation"]["valid"])
        self.assertTrue(result["preview_only"])

    @patch("medical_report_mcp_server.validate_manifest")
    def test_word_generation_uses_confirmed_format_authority_and_stays_non_delivery(
        self, validate_manifest
    ) -> None:
        template = self.root / "template.docx"
        document = Document()
        if "Body Text First Indent" not in {style.name for style in document.styles}:
            document.styles.add_style("Body Text First Indent", WD_STYLE_TYPE.PARAGRAPH)
        document.add_paragraph("旧模板正文应清除")
        document.save(template)
        config = self.root / "word-format-authority.json"
        config.write_text(
            json.dumps(
                {
                    "schema_version": "1.0",
                    "status": "confirmed",
                    "profile_id": "PROFILE-MCP-TEST",
                    "confirmed_by": "tester",
                    "confirmed_at": "2026-08-24T12:00:00+08:00",
                    "authority": {
                        "template_path": template.name,
                        "template_sha256": sha256_file(template),
                    },
                    "semantic_styles": {
                        "heading_1": "Heading 1",
                        "heading_2": "Heading 2",
                        "heading_3": "Heading 3",
                        "heading_4": "Heading 4",
                        "heading_5": "Heading 5",
                        "heading_6": "Heading 6",
                        "heading_7": "Heading 7",
                        "body": "Body Text First Indent",
                        "table": "Table Grid",
                        "table_header": "Normal",
                        "table_body": "Normal",
                    },
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        manifest = {
            "project_code": "TEST-001",
            "manifest_id": "MANIFEST-001",
            "manifest_hash": "manifest-hash",
            "package_id": "PACK-001",
            "package_content_hash": "package-hash",
            "status": "complete",
            "preview_only": False,
            "construction_list_import": {"display_payload": {"sheets": []}},
            "application_software_solution": {
                "items": [
                    {
                        "original_name": "主数据管理",
                        "status": "verbatim",
                        "fragments": [
                            {
                                "relative_heading_path": [],
                                "clean_text": "标准方案正文。",
                                "source_section_id": "SECTION-001",
                            }
                        ],
                    }
                ]
            },
        }
        validation = {
            "valid": True,
            "manifest_id": "MANIFEST-001",
            "manifest_hash": "manifest-hash",
            "package_id": "PACK-001",
            "package_content_hash": "package-hash",
            "validation_hash": "validation-hash",
        }
        validate_manifest.return_value = validation
        manifest_path = self.root / "construction-assembly-manifest.json"
        validation_path = self.root / "construction-assembly-validation.json"
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
        validation_path.write_text(json.dumps(validation, ensure_ascii=False), encoding="utf-8")
        markdown = self.root / "solution.md"
        markdown.write_text(
            "## 软件建设清单\n\n\n## 应用软件建设方案\n\n### 主数据管理\n\n标准方案正文。\n",
            encoding="utf-8",
        )
        output = self.root / "solution.docx"

        result = self.server.word_generate(
            {
                "project_root": str(self.root),
                "input_markdown_path": str(markdown),
                "output_docx_path": str(output),
                "project_name": "MCP测试项目",
                "format_config_path": str(config),
                "database_path": str(self.database),
                "project_code": "TEST-001",
                "assembly_manifest_path": str(manifest_path),
                "assembly_validation_path": str(validation_path),
            }
        )

        self.assertEqual(result["status"], "structure_pass_render_required")
        self.assertTrue(result["structural_validation_passed"])
        self.assertEqual(result["visual_render_review"], "not_run")
        self.assertFalse(result["delivery_ready"])
        self.assertEqual(result["format_profile_id"], "PROFILE-MCP-TEST")
        self.assertTrue(result["fragment_wrapper_applied"])
        self.assertEqual(result["heading_levels"]["1"], 1)
        self.assertEqual(result["heading_levels"]["2"], 2)
        self.assertEqual(result["heading_levels"]["3"], 1)
        generated = Document(output)
        self.assertTrue(
            any(paragraph.text == "标准方案正文。" for paragraph in generated.paragraphs)
        )
        self.assertTrue(output.is_file())

        markdown.write_text(
            markdown.read_text(encoding="utf-8") + "\nAI自行补写。\n",
            encoding="utf-8",
        )
        with self.assertRaisesRegex(ValueError, "exact output"):
            self.server.word_generate(
                {
                    "project_root": str(self.root),
                    "input_markdown_path": str(markdown),
                    "output_docx_path": str(self.root / "tampered.docx"),
                    "project_name": "MCP测试项目",
                    "format_config_path": str(config),
                    "database_path": str(self.database),
                    "project_code": "TEST-001",
                    "assembly_manifest_path": str(manifest_path),
                    "assembly_validation_path": str(validation_path),
                }
            )

        blocked = self.server.dispatch(
            {
                "jsonrpc": "2.0",
                "id": 5,
                "method": "tools/call",
                "params": {
                    "name": "word_generate",
                    "arguments": {
                        "project_root": str(self.root),
                        "input_markdown_path": str(markdown),
                        "output_docx_path": str(self.root / "should-not-exist.docx"),
                        "project_name": "MCP测试项目",
                        "format_config_path": str(config),
                        "markdown_mode": "full_report",
                    },
                },
            }
        )
        self.assertTrue(blocked["result"]["isError"])
        self.assertIn("第1章", blocked["result"]["structuredContent"]["message"])


if __name__ == "__main__":
    unittest.main()
